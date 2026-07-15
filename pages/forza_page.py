import re
import random
import os
from pathlib import Path
import pyodbc
import openpyxl
from openpyxl.utils import column_index_from_string
import allure
from playwright.sync_api import Page, expect

class ForzaPage:
    def __init__(self, page: Page, db_connection_string: str = ""):
        self.page = page
        self.db_connection_string = db_connection_string #
        
        # Variables de estado
        self.ruta_excel: str = ""
        self.hoja: str = ""
        self.columna: str = ""
        self.url_actual: str = ""
        self.token: str = ""
        self.entorno: str = ""
        self.guia_generada: str = ""
        self.pais_actual: str = ""

    def _take_screenshot(self, name: str = "screenshot"):
        try:
            screenshot = self.page.screenshot(type="png")
            allure.attach(screenshot, name=name, attachment_type=allure.attachment_type.PNG)
        except Exception as e:
            print(f"Error taking screenshot: {e}")
    
    @allure.step("Navegar a la URL: '{url}'")
    def go_to_page_web(self, url: str):
        self.url_actual = url
        ultimo_error = None
        for intento in range(3):
            try:
                self.page.goto(url, timeout=120000, wait_until="load")
                break
            except Exception as error:
                ultimo_error = error
                if intento == 2:
                    raise
                self.page.wait_for_timeout(5000)
        self._take_screenshot("navigation")

    @allure.step("Validar que el título sea: '{titulo_esperado}'")
    def assert_title(self, titulo_esperado: str):
        assert self.page.title() == titulo_esperado, f"Error: Título esperado '{titulo_esperado}', actual '{self.page.title()}'"
        self._take_screenshot("title_validation")

    @allure.step("Seleccionar país: '{country}'")
    def select_country(self, country: str):
        self.pais_actual = country
        patron_regex = re.compile(f"^{re.escape(country)}$")
        test_ids_por_pais = {
            "guatemala": "pw-login-option-country-gt",
            "honduras": "pw-login-option-country-hn",
            "el salvador": "pw-login-option-country-sv",
        }
        country_test_id = test_ids_por_pais.get(country.strip().lower())
        opciones_pais = [
            self.page.locator(f"[data-testid='{country_test_id}']") if country_test_id else self.page.locator("not-found"),
            self.page.get_by_role("button", name=re.compile(country, re.IGNORECASE)),
            self.page.locator("ion-item").filter(has_text=patron_regex),
            self.page.locator(".container-icon").filter(has_text=patron_regex),
            self.page.locator("div").filter(has_text=patron_regex).get_by_role("emphasis"),
        ]

        def login_visible():
            usuario_corp = self.page.get_by_text("Usuario Corporativo", exact=True)
            correo = self.page.get_by_text("Correo electrónico", exact=True)
            return (
                usuario_corp.count() > 0 and usuario_corp.first.is_visible()
            ) or (
                correo.count() > 0 and correo.first.is_visible()
            )

        for opcion in opciones_pais:
            try:
                if opcion.count() == 0:
                    continue
                opcion.first.click(timeout=10000, force=True)
                try:
                    self.page.wait_for_load_state("networkidle", timeout=10000)
                except Exception:
                    pass
                self.page.wait_for_timeout(1000)
                if login_visible():
                    break
            except Exception:
                continue
        else:
            raise AssertionError(f"No se pudo seleccionar el país: {country}")

        self.page.wait_for_load_state("load")
        self._take_screenshot("country_selection")

    @allure.step("Seleccionar país en App Courier: '{country}'")
    def select_country_courier(self, country: str):
        country_text = self.page.get_by_text(country, exact=True)
        if country_text.count() > 0 and country_text.first.is_visible():
            self._take_screenshot("courier_country_selection")
            return

        try:
            self.page.get_by_role("combobox").select_option(label=country, timeout=10000)
        except Exception as first_error:
            try:
                self.page.locator("select").first.select_option(label=country, timeout=5000)
            except Exception:
                try:
                    self.page.get_by_role("combobox").first.click(timeout=5000)
                    self.page.get_by_text(country, exact=True).click(timeout=5000)
                except Exception:
                    if country_text.count() > 0:
                        return
                    raise first_error
        self._take_screenshot("courier_country_selection")

    # ==============================================================================
    # FLUJOS DE LOGIN
    # ==============================================================================

    @allure.step("Login Portal Web - Correo: '{correo}'")
    def login(self, correo: str, passw: str):
        self.page.get_by_role("textbox", name="Ingresar Usuario").click()
        self.page.get_by_role("textbox", name="Ingresar Usuario").fill(correo)
        
        self.page.get_by_role("textbox", name="Ingresar Contraseña").click()
        self.page.get_by_role("textbox", name="Ingresar Contraseña").fill(passw)
        
        self.page.get_by_role("button", name="Iniciar sesión").click()
        self.page.wait_for_url(f"**/design/dashboard", wait_until="networkidle")
        self._take_screenshot("login_success")

    @allure.step("Abrir Tienda Usuario")
    def abrir_tienda_usuario(self, url: str):
        self.url_actual = url
        self.page.goto(url, timeout=120000, wait_until="load")
        try:
            self.page.wait_for_load_state("networkidle", timeout=30000)
        except Exception:
            pass
        self._seleccionar_pais_tienda_si_aplica("Guatemala")
        expect(self.page.get_by_text(re.compile(r"Detalle Carrito|Iniciar Sesi[oó]n|Membres", re.IGNORECASE)).first).to_be_visible(timeout=60000)
        self._take_screenshot("tienda_home")

    def _seleccionar_pais_tienda_si_aplica(self, pais: str):
        selector_pais_visible = self.page.get_by_text(
            re.compile(r"Para comenzar, por favor selecciona tu pa[ií]s", re.IGNORECASE)
        ).first
        try:
            if selector_pais_visible.count() == 0 or not selector_pais_visible.is_visible():
                return
        except Exception:
            return

        opciones_pais = [
            self.page.locator(".container-icon").filter(has_text=re.compile(pais, re.IGNORECASE)).first,
            self.page.get_by_text(pais, exact=True).first.locator("xpath=ancestor::*[contains(@class,'container-icon')][1]"),
            self.page.get_by_text(pais, exact=True).first.locator("xpath=.."),
            self.page.get_by_role("button", name=re.compile(pais, re.IGNORECASE)).first,
            self.page.locator(".btn-country").filter(has_text=re.compile(pais, re.IGNORECASE)).first,
            self.page.locator("button, div, a").filter(has_text=re.compile(rf"^{re.escape(pais)}$", re.IGNORECASE)).first,
            self.page.get_by_text(pais, exact=True).first,
        ]
        ultimo_error = None
        for opcion in opciones_pais:
            try:
                if opcion.count() == 0:
                    continue
                opcion.scroll_into_view_if_needed(timeout=10000)
                opcion.click(timeout=30000, force=True)
                self.page.wait_for_timeout(1500)
                if not selector_pais_visible.is_visible():
                    break
            except Exception as error:
                ultimo_error = error
                continue
        else:
            try:
                self.page.get_by_text(pais, exact=True).first.evaluate(
                    """(el) => {
                        const tarjeta = el.closest('.container-icon') || el.parentElement;
                        tarjeta.click();
                    }"""
                )
                self.page.wait_for_timeout(1500)
            except Exception as error:
                raise AssertionError(f"No se pudo seleccionar el país de tienda: {pais}. Error: {ultimo_error or error}")

        try:
            if selector_pais_visible.is_visible():
                self.page.get_by_text(pais, exact=True).first.evaluate(
                    """(el) => {
                        const tarjeta = el.closest('.container-icon') || el.parentElement;
                        tarjeta.click();
                    }"""
                )
                self.page.wait_for_timeout(1500)
        except Exception:
            pass

        try:
            self.page.wait_for_load_state("networkidle", timeout=30000)
        except Exception:
            pass
        expect(self.page.get_by_text(re.compile(r"Detalle Carrito|Iniciar Sesi[oó]n|Membres", re.IGNORECASE)).first).to_be_visible(timeout=60000)
        self.page.wait_for_timeout(1000)

    @allure.step("Login Tienda Usuario - Correo: '{correo}'")
    def login_tienda_usuario(self, correo: str, passw: str):
        if self._tienda_usuario_autenticado():
            self._take_screenshot("tienda_login_success")
            return
        self._abrir_login_tienda()
        self._fill_tienda_textbox(re.compile(r"Nombre de Usuario|Usuario|Correo|Email", re.IGNORECASE), correo)
        self._fill_tienda_textbox(re.compile(r"Contrase", re.IGNORECASE), passw)
        self._submit_login_tienda()
        self._validar_tienda_autenticada()
        self._take_screenshot("tienda_login_success")

    def _tienda_usuario_autenticado(self) -> bool:
        try:
            if self._texto_visible_tienda(r"Inicia(?:r)? sesi[oó]n"):
                return False
        except Exception:
            pass

        try:
            return not self._tienda_login_visible()
        except Exception:
            return True

    def _submit_login_tienda(self):
        try:
            hizo_click = self.page.evaluate(
                """() => {
                    const normalizar = (valor) => (valor || '').replace(/\\s+/g, ' ').trim().toLowerCase();
                    const visible = (item) => {
                        const estilo = window.getComputedStyle(item);
                        const caja = item.getBoundingClientRect();
                        return estilo.visibility !== 'hidden' && estilo.display !== 'none' && caja.width > 0 && caja.height > 0;
                    };
                    const opciones = Array.from(document.querySelectorAll('button, input[type="submit"], a'))
                        .filter(visible)
                        .filter((item) => !item.closest('nav, header'))
                        .filter((item) => normalizar(item.innerText || item.value || item.textContent).match(/^iniciar sesi[oó]n$/i));
                    const boton = opciones.sort((a, b) => b.getBoundingClientRect().top - a.getBoundingClientRect().top)[0];
                    if (!boton) {
                        return false;
                    }
                    boton.scrollIntoView({ block: 'center', inline: 'center' });
                    boton.click();
                    return true;
                }"""
            )
            if hizo_click:
                return
        except Exception:
            pass

        self._click_tienda_por_texto(r"^Iniciar sesi[oó]n$", timeout=30000)

    def _validar_tienda_autenticada(self):
        try:
            self.page.wait_for_load_state("networkidle", timeout=30000)
        except Exception:
            pass
        self.page.wait_for_function(
            """() => {
                const normalizar = (valor) => (valor || '').replace(/\\s+/g, ' ').trim().toLowerCase();
                const visible = (item) => {
                    const estilo = window.getComputedStyle(item);
                    const caja = item.getBoundingClientRect();
                    return estilo.visibility !== 'hidden' && estilo.display !== 'none' && caja.width > 0 && caja.height > 0;
                };
                const textos = Array.from(document.querySelectorAll('button, a, span, div, p'))
                    .filter(visible)
                    .map((item) => normalizar(item.innerText || item.textContent))
                    .filter(Boolean);
                return !textos.some((texto) => /^iniciar sesi[oó]n$/i.test(texto));
            }""",
            timeout=120000,
        )

    def _texto_visible_tienda(self, texto_regex: str) -> bool:
        return bool(
            self.page.evaluate(
                """(textoRegex) => {
                    const patron = new RegExp(textoRegex, 'i');
                    const visible = (item) => {
                        const estilo = window.getComputedStyle(item);
                        const caja = item.getBoundingClientRect();
                        return estilo.visibility !== 'hidden' && estilo.display !== 'none' && caja.width > 0 && caja.height > 0;
                    };
                    return Array.from(document.querySelectorAll('button, a, span, div, p'))
                        .filter(visible)
                        .some((item) => patron.test((item.innerText || item.textContent || '').replace(/\\s+/g, ' ').trim()));
                }""",
                texto_regex,
            )
        )

    def _abrir_login_tienda(self):
        if self._tienda_usuario_autenticado():
            return
        if self._tienda_login_visible():
            return

        try:
            match = re.match(r"^(https?://[^/]+)", self.page.url or self.url_actual)
            if match:
                self.page.goto(f"{match.group(1)}/inicio-sesion", timeout=120000, wait_until="load")
                expect(self.page.get_by_text(re.compile(r"Nombre de Usuario|Contrase", re.IGNORECASE)).first).to_be_visible(timeout=30000)
                return
        except Exception:
            pass

        try:
            self.page.get_by_text("Iniciar Sesión", exact=False).first.click(timeout=10000)
            self.page.wait_for_timeout(800)
            self.page.locator("a.dropdown-item").filter(has_text="Iniciar Sesión").first.click(timeout=10000)
            expect(self.page.get_by_text(re.compile(r"Nombre de Usuario|Contrase", re.IGNORECASE)).first).to_be_visible(timeout=30000)
            return
        except Exception:
            pass

        patron_login = re.compile(r"Inicia(?:r)? sesi[oó]n|Login", re.IGNORECASE)
        opciones_menu = [
            self.page.locator("a.nav-link, button, a").filter(has_text=patron_login).first,
            self.page.get_by_text(patron_login).first,
        ]
        for opcion in opciones_menu:
            try:
                if opcion.count() == 0:
                    continue
                opcion.click(timeout=10000)
                self.page.wait_for_timeout(800)
                if self._tienda_login_visible():
                    return
                opcion_dropdown = self.page.locator(".dropdown-menu a, .dropdown-item, a").filter(has_text=patron_login).last
                if opcion_dropdown.count() > 0 and opcion_dropdown.is_visible():
                    opcion_dropdown.click(timeout=10000)
                    self.page.wait_for_timeout(1000)
                    if self._tienda_login_visible():
                        return
            except Exception:
                continue

        self._click_tienda_por_texto(r"Inicia(?:r)? sesi[oó]n", timeout=15000)
        self.page.wait_for_timeout(800)
        try:
            self.page.locator(".dropdown-menu a, .dropdown-item, a").filter(has_text=patron_login).last.click(timeout=10000)
        except Exception:
            pass
        expect(self.page.get_by_text(re.compile(r"Nombre de Usuario|Contrase", re.IGNORECASE)).first).to_be_visible(timeout=30000)

    def _tienda_login_visible(self) -> bool:
        try:
            campo_usuario = self.page.get_by_text(re.compile(r"Nombre de Usuario|Contrase", re.IGNORECASE)).first
            return campo_usuario.count() > 0 and campo_usuario.is_visible()
        except Exception:
            return False

    def _fill_tienda_textbox(self, label_regex, value: str):
        opciones = [
            self.page.get_by_role("textbox", name=label_regex).first,
            self.page.get_by_placeholder(label_regex).first,
            self.page.locator("input").filter(has=self.page.locator("not-found")),
        ]

        for opcion in opciones[:2]:
            try:
                if opcion.count() == 0:
                    continue
                opcion.click(timeout=10000)
                self.page.keyboard.press("Control+A")
                self.page.keyboard.type(value)
                return
            except Exception:
                continue

        inputs = self.page.locator("input")
        input_count = inputs.count()
        for index in range(input_count):
            campo = inputs.nth(index)
            try:
                tipo = (campo.get_attribute("type") or "").lower()
                placeholder = campo.get_attribute("placeholder") or ""
                aria = campo.get_attribute("aria-label") or ""
                texto_campo = f"{placeholder} {aria}"
                if label_regex.search(texto_campo) or ("password" in tipo and label_regex.search("Contraseña")):
                    campo.click(timeout=5000)
                    self.page.keyboard.press("Control+A")
                    self.page.keyboard.type(value)
                    return
            except Exception:
                continue

        raise AssertionError(f"No se encontró el campo de tienda para: {label_regex.pattern}")

    @allure.step("Comprar membresia '{producto}' en Tienda Usuario")
    def comprar_membresia_tienda(self, producto: str, facturacion: str, metodo_pago: str):
        self._click_tienda_por_texto(r"Membres[ií]as", timeout=60000)
        expect(self.page.get_by_text(producto, exact=True).first).to_be_visible(timeout=60000)
        self._seleccionar_producto_tienda(producto)
        self._take_screenshot("tienda_producto")

        if not self._carrito_tienda_tiene_productos():
            self._click_tienda_por_texto(r"Agregar (?:al|a) carrito", timeout=60000)
        self._ir_al_carrito_tienda()
        expect(self.page.get_by_text(re.compile(r"Carrito de Compras|Informaci[oó]n de Facturaci[oó]n", re.IGNORECASE)).first).to_be_visible(timeout=60000)
        self._take_screenshot("tienda_carrito")

        self._seleccionar_opcion_tienda(facturacion)
        self._seleccionar_opcion_tienda(metodo_pago)
        self._aceptar_terminos_tienda()
        self._click_tienda_por_texto(r"Pagar Carrito", timeout=30000)
        self._confirmar_datos_tienda()

        expect(self.page.get_by_text(re.compile(r"Tu solicitud fue operada exitosamente|Siguiente", re.IGNORECASE)).first).to_be_visible(timeout=240000)
        self._click_tienda_por_texto(r"Siguiente", timeout=240000)
        self._take_screenshot("tienda_transaccion_exitosa")

    def _seleccionar_producto_tienda(self, producto: str):
        try:
            self.page.wait_for_function(
                "() => !document.body.innerText.includes('Cargando...')",
                timeout=30000,
            )
        except Exception:
            pass

        producto_regex = re.compile(re.escape(producto), re.IGNORECASE)
        tarjeta_producto = self.page.get_by_text(producto, exact=True).first.locator(
            "xpath=ancestor::*[contains(@class,'col-')][1]"
        )
        opciones = [
            tarjeta_producto.locator("img.cursor-pointer, img").first,
            tarjeta_producto,
            self.page.locator("[class*='card']").filter(has_text=producto_regex).first,
            self.page.locator("[class*='product']").filter(has_text=producto_regex).first,
            self.page.get_by_text(producto, exact=True).first.locator("xpath=ancestor::*[contains(@class,'card') or contains(@class,'product')][1]"),
            self.page.get_by_text(producto, exact=True).first,
        ]

        ultimo_error = None
        for opcion in opciones:
            try:
                if opcion.count() == 0:
                    continue
                opcion.scroll_into_view_if_needed(timeout=10000)
                opcion.click(timeout=30000, force=True)
                self.page.wait_for_timeout(1000)
                if "/detalle-producto/" not in self.page.url:
                    continue
                expect(self.page.get_by_text(re.compile(r"Agregar (?:al|a) carrito", re.IGNORECASE)).first).to_be_visible(timeout=60000)
                return
            except Exception as error:
                ultimo_error = error
                continue

        if producto.strip().lower() == "club forza":
            try:
                match = re.match(r"^(https?://[^/]+)", self.page.url or self.url_actual)
                if match:
                    self.page.goto(f"{match.group(1)}/GT/detalle-producto/club-forza", timeout=120000, wait_until="load")
                    expect(self.page.get_by_text(re.compile(r"Agregar (?:al|a) carrito", re.IGNORECASE)).first).to_be_visible(timeout=60000)
                    return
            except Exception as error:
                ultimo_error = error

        raise AssertionError(f"No se pudo seleccionar el producto de tienda: {producto}. Error: {ultimo_error}")

    def _ir_al_carrito_tienda(self):
        try:
            self._click_tienda_por_texto(r"Ir al carrito", timeout=15000)
        except AssertionError:
            if not self._abrir_carrito_tienda_por_icono():
                self._abrir_carrito_tienda_por_url()

        try:
            self.page.wait_for_load_state("networkidle", timeout=30000)
        except Exception:
            pass
        self._esperar_loader_tienda(timeout=120000)
        if not self._texto_visible_tienda(r"Carrito de Compras|Informaci[oó]n de Facturaci[oó]n|Detalle de la compra|Total a pagar"):
            if not self._abrir_carrito_tienda_por_icono():
                self._abrir_carrito_tienda_por_url()
            self._esperar_loader_tienda(timeout=120000)
        expect(self.page.get_by_text(re.compile(r"Carrito de Compras|Informaci[oó]n de Facturaci[oó]n|Detalle de la compra|Total a pagar", re.IGNORECASE)).first).to_be_visible(timeout=120000)

    def _abrir_carrito_tienda_por_icono(self) -> bool:
        try:
            abierto = self.page.evaluate(
                """() => {
                    const visible = (item) => {
                        const estilo = window.getComputedStyle(item);
                        const caja = item.getBoundingClientRect();
                        return estilo.visibility !== 'hidden' && estilo.display !== 'none' && caja.width > 0 && caja.height > 0;
                    };
                    const candidatos = Array.from(document.querySelectorAll('img, i, em, svg, span'))
                        .filter(visible)
                        .filter((item) => {
                            const src = item.currentSrc || item.src || '';
                            const clase = item.className || '';
                            const texto = item.innerText || item.textContent || '';
                            return /cart|shopping|carrito|numberIcon/i.test(`${src} ${clase} ${texto}`);
                        });
                    for (const candidato of candidatos) {
                        const clickeable = candidato.closest('a, button, [role="button"], .cursor-pointer, div') || candidato;
                        if (visible(clickeable)) {
                            clickeable.scrollIntoView({ block: 'center', inline: 'center' });
                            clickeable.click();
                            return true;
                        }
                    }
                    const elementos = Array.from(document.elementsFromPoint(window.innerWidth - 235, 50));
                    const objetivo = elementos.find((item) => visible(item) && /cart|shopping|carrito|numberIcon/i.test(`${item.className || ''} ${item.innerText || ''} ${item.textContent || ''}`));
                    if (objetivo) {
                        (objetivo.closest('a, button, [role="button"], .cursor-pointer, div') || objetivo).click();
                        return true;
                    }
                    return false;
                }"""
            )
            if abierto:
                self.page.wait_for_timeout(1500)
                return True
        except Exception:
            pass
        return False

    def _abrir_carrito_tienda_por_url(self):
        match = re.match(r"^(https?://[^/]+)", self.page.url or self.url_actual)
        base_url = match.group(1) if match else "https://qa-tienda.forzadeliveryexpress.com"
        self.page.goto(f"{base_url}/GT/carrito", timeout=120000, wait_until="domcontentloaded")
        try:
            self.page.wait_for_load_state("load", timeout=30000)
        except Exception:
            pass

    def _carrito_tienda_tiene_productos(self) -> bool:
        try:
            return bool(
                self.page.evaluate(
                    """() => {
                        const texto = (document.body.innerText || '').replace(/\\s+/g, ' ');
                        const contador = Array.from(document.querySelectorAll('.numberIcon, [class*="numberIcon"], sup, span'))
                            .map((item) => (item.innerText || item.textContent || '').trim())
                            .find((valor) => /^\\d+$/.test(valor));
                        return Number(contador || 0) > 0 || /Cantidad de productos\\s+[1-9]/i.test(texto);
                    }"""
                )
            )
        except Exception:
            return False

    def _esperar_loader_tienda(self, timeout: int = 60000):
        try:
            self.page.wait_for_function(
                """() => {
                    const texto = (document.body.innerText || '').replace(/\\s+/g, ' ');
                    const loaderVisible = Array.from(document.querySelectorAll('.swal2-container, .hwa-spinner, .preloader'))
                        .some((item) => {
                            const estilo = window.getComputedStyle(item);
                            const caja = item.getBoundingClientRect();
                            return estilo.visibility !== 'hidden'
                                && estilo.display !== 'none'
                                && caja.width > 0
                                && caja.height > 0
                                && /Cargando/i.test(item.innerText || item.textContent || texto);
                        });
                    return !loaderVisible;
                }""",
                timeout=timeout,
            )
        except Exception:
            try:
                self.page.keyboard.press("Escape")
                self.page.evaluate(
                    """() => {
                        document.querySelectorAll('.swal2-container').forEach((item) => item.remove());
                        document.querySelectorAll('.preloader, .preloader-hidden, .hwa-spinner').forEach((item) => item.remove());
                        document.body.classList.remove('swal2-shown', 'swal2-height-auto');
                        document.querySelectorAll('[aria-hidden="true"]').forEach((item) => item.removeAttribute('aria-hidden'));
                    }"""
                )
                self.page.wait_for_timeout(500)
            except Exception:
                pass

    def _seleccionar_opcion_tienda(self, texto: str):
        opcion_regex = re.compile(rf"(^|\s){re.escape(texto)}($|\s)", re.IGNORECASE)
        opciones = [
            self.page.get_by_label(opcion_regex).first,
            self.page.locator("label, mat-radio-button, .mat-radio-button, .form-check, ion-radio").filter(has_text=opcion_regex).first,
        ]
        for opcion in opciones:
            try:
                if opcion.count() == 0:
                    continue
                opcion.scroll_into_view_if_needed(timeout=10000)
                opcion.click(timeout=10000, force=True)
                self.page.wait_for_timeout(500)
                return
            except Exception:
                continue

        try:
            seleccionado = self.page.evaluate(
                """(texto) => {
                    const normalizar = (valor) => (valor || '').replace(/\\s+/g, ' ').trim().toLowerCase();
                    const esperado = normalizar(texto);
                    const visible = (item) => {
                        const estilo = window.getComputedStyle(item);
                        const caja = item.getBoundingClientRect();
                        return estilo.visibility !== 'hidden' && estilo.display !== 'none' && caja.width > 0 && caja.height > 0;
                    };
                    const nodos = Array.from(document.querySelectorAll('label, mat-radio-button, .mat-radio-button, .form-check, ion-radio, p, span, div'))
                        .filter(visible)
                        .map((item) => ({ item, contenido: normalizar(item.innerText || item.textContent) }))
                        .filter(({ contenido }) => contenido === esperado || contenido.startsWith(`${esperado} `))
                        .sort((a, b) => a.contenido.length - b.contenido.length);
                    const nodo = nodos.find(({ contenido }) => contenido.length <= esperado.length + 24);
                    if (!nodo) {
                        return false;
                    }
                    const item = nodo.item;
                    const contenedor = item.closest('label, mat-radio-button, .mat-radio-button, .form-check, ion-radio');
                    const input = contenedor ? contenedor.querySelector('input[type="radio"], input[type="checkbox"]') : null;
                    const clickeable = input || contenedor || item;
                    clickeable.scrollIntoView({ block: 'center', inline: 'center' });
                    clickeable.click();
                    return true;
                }""",
                texto,
            )
            if seleccionado:
                self.page.wait_for_timeout(500)
                return
        except Exception:
            pass

        try:
            seleccionado = self.page.evaluate(
                """(texto) => {
                    const normalizar = (valor) => (valor || '').replace(/\\s+/g, ' ').trim().toLowerCase();
                    const esperado = normalizar(texto);
                    const nodos = Array.from(document.querySelectorAll('p, span, div')).filter((item) => {
                        const contenido = normalizar(item.innerText || item.textContent);
                        return contenido === esperado;
                    });
                    if (nodos.length === 0) {
                        return false;
                    }
                    const nodo = nodos.sort((a, b) => {
                        const ac = a.getBoundingClientRect();
                        const bc = b.getBoundingClientRect();
                        return (ac.width * ac.height) - (bc.width * bc.height);
                    })[0];
                    const caja = nodo.getBoundingClientRect();
                    const x = Math.max(caja.left - 18, 0);
                    const y = caja.top + caja.height / 2;
                    const radioCercano = document.elementFromPoint(x, y);
                    const clickeable = radioCercano || nodo;
                    clickeable.scrollIntoView({ block: 'center', inline: 'center' });
                    clickeable.click();
                    return true;
                }""",
                texto,
            )
            if seleccionado:
                self.page.wait_for_timeout(500)
                return
        except Exception:
            pass

        raise AssertionError(f"No se pudo seleccionar la opción de tienda: {texto}")

    def _aceptar_terminos_tienda(self):
        self._click_tienda_por_texto(r"Estoy de acuerdo con los t[eé]rminos y condiciones", timeout=30000)
        expect(self.page.get_by_text(re.compile(r"T[eé]rminos y condiciones", re.IGNORECASE)).first).to_be_visible(timeout=30000)
        boton_aceptar = self.page.get_by_role("button", name=re.compile(r"Aceptar", re.IGNORECASE)).first
        for _ in range(12):
            try:
                if boton_aceptar.count() > 0 and boton_aceptar.is_visible():
                    boton_aceptar.click(timeout=5000)
                    self.page.wait_for_timeout(800)
                    if not self._modal_terminos_visible():
                        return
            except Exception:
                pass
            try:
                self.page.evaluate(
                    """() => {
                        const visibles = Array.from(document.querySelectorAll('div, section, article'))
                            .filter((item) => {
                                const estilo = window.getComputedStyle(item);
                                const caja = item.getBoundingClientRect();
                                return estilo.visibility !== 'hidden'
                                    && estilo.display !== 'none'
                                    && caja.width > 250
                                    && caja.height > 250
                                    && item.scrollHeight > item.clientHeight + 50;
                            })
                            .sort((a, b) => (b.scrollHeight - b.clientHeight) - (a.scrollHeight - a.clientHeight));
                        const modalScroll = visibles[0];
                        if (modalScroll) {
                            modalScroll.scrollTop = modalScroll.scrollTop + modalScroll.clientHeight;
                        }
                    }"""
                )
            except Exception:
                self.page.mouse.wheel(0, 900)
            self.page.wait_for_timeout(300)

        try:
            aceptado = self.page.evaluate(
                """() => {
                    const normalizar = (valor) => (valor || '').replace(/\\s+/g, ' ').trim().toLowerCase();
                    const visible = (item) => {
                        const estilo = window.getComputedStyle(item);
                        const caja = item.getBoundingClientRect();
                        return estilo.visibility !== 'hidden' && estilo.display !== 'none' && caja.width > 0 && caja.height > 0;
                    };

                    for (let i = 0; i < 8; i += 1) {
                        Array.from(document.querySelectorAll('div, section, article'))
                            .filter(visible)
                            .filter((item) => item.scrollHeight > item.clientHeight + 20)
                            .forEach((item) => { item.scrollTop = item.scrollHeight; });
                    }

                    const opciones = Array.from(document.querySelectorAll('button, a, span, label, div'))
                        .filter(visible)
                        .filter((item) => normalizar(item.innerText || item.textContent) === 'aceptar')
                        .sort((a, b) => b.getBoundingClientRect().top - a.getBoundingClientRect().top);
                    if (opciones.length > 0) {
                        opciones[0].click();
                        return true;
                    }

                    const modal = Array.from(document.querySelectorAll('.modal-content, .swal2-popup, div'))
                        .filter(visible)
                        .filter((item) => /t[eé]rminos y condiciones/i.test(item.innerText || item.textContent || ''))
                        .sort((a, b) => (b.getBoundingClientRect().width * b.getBoundingClientRect().height) - (a.getBoundingClientRect().width * a.getBoundingClientRect().height))[0];
                    if (!modal) {
                        return false;
                    }
                    const caja = modal.getBoundingClientRect();
                    const elemento = document.elementFromPoint(caja.left + caja.width / 2, caja.bottom - 35);
                    if (elemento) {
                        elemento.click();
                        return true;
                    }
                    return false;
                }"""
            )
            if aceptado:
                self.page.wait_for_timeout(800)
                if not self._modal_terminos_visible():
                    return
        except Exception:
            pass

        try:
            cerrado = self.page.evaluate(
                """() => {
                    const modal = document.querySelector('#terminosCondiciones');
                    if (!modal) {
                        return true;
                    }
                    const scrollables = Array.from(modal.querySelectorAll('div, section, article'))
                        .filter((item) => item.scrollHeight > item.clientHeight + 20);
                    scrollables.forEach((item) => { item.scrollTop = item.scrollHeight; });
                    const visible = (item) => {
                        const estilo = window.getComputedStyle(item);
                        const caja = item.getBoundingClientRect();
                        return estilo.visibility !== 'hidden' && estilo.display !== 'none' && caja.width > 0 && caja.height > 0;
                    };
                    const botones = Array.from(modal.querySelectorAll('button, a, span, label, div'))
                        .filter(visible)
                        .filter((item) => /aceptar/i.test(item.innerText || item.textContent || ''));
                    const objetivo = botones[botones.length - 1];
                    if (objetivo) {
                        objetivo.scrollIntoView({ block: 'center', inline: 'center' });
                        objetivo.click();
                        return true;
                    }
                    const caja = modal.getBoundingClientRect();
                    const elemento = document.elementFromPoint(caja.left + caja.width / 2, caja.bottom - 40);
                    if (elemento) {
                        elemento.click();
                        return true;
                    }
                    return false;
                }"""
            )
            if cerrado:
                self.page.wait_for_timeout(1000)
                if not self._modal_terminos_visible():
                    return
        except Exception:
            pass

        boton_aceptar.click(timeout=10000, force=True)
        expect(self.page.locator("#terminosCondiciones")).not_to_be_visible(timeout=30000)

    def _modal_terminos_visible(self) -> bool:
        try:
            modal = self.page.locator("#terminosCondiciones").first
            return modal.count() > 0 and modal.is_visible()
        except Exception:
            return False

    def _confirmar_datos_tienda(self):
        expect(self.page.get_by_text(re.compile(r"Confirmar datos", re.IGNORECASE)).first).to_be_visible(timeout=60000)
        self._take_screenshot("tienda_confirmar_datos")
        try:
            boton_texto = self.page.get_by_text("Confirmar", exact=True).last
            boton_texto.scroll_into_view_if_needed(timeout=10000)
            boton_texto.click(timeout=10000, force=True)
            self.page.wait_for_timeout(1500)
            if not self._texto_visible_tienda(r"^Confirmar datos$"):
                return
        except Exception:
            pass

        try:
            confirmado = self.page.evaluate(
                """() => {
                    const normalizar = (valor) => (valor || '').replace(/\\s+/g, ' ').trim().toLowerCase();
                    const visible = (item) => {
                        const estilo = window.getComputedStyle(item);
                        const caja = item.getBoundingClientRect();
                        return estilo.visibility !== 'hidden' && estilo.display !== 'none' && caja.width > 0 && caja.height > 0;
                    };
                    const modales = Array.from(document.querySelectorAll('.modal-content, app-modal, div'))
                        .filter(visible)
                        .filter((item) => /confirmar datos/i.test(item.innerText || item.textContent || ''))
                        .sort((a, b) => (b.getBoundingClientRect().width * b.getBoundingClientRect().height) - (a.getBoundingClientRect().width * a.getBoundingClientRect().height));
                    const modal = modales[0] || document.body;
                    const boton = Array.from(modal.querySelectorAll('button, a, span, label, div'))
                        .filter(visible)
                        .filter((item) => normalizar(item.innerText || item.textContent) === 'confirmar')
                        .sort((a, b) => b.getBoundingClientRect().top - a.getBoundingClientRect().top)[0];
                    if (boton) {
                        boton.scrollIntoView({ block: 'center', inline: 'center' });
                        boton.click();
                        return true;
                    }

                    const caja = modal.getBoundingClientRect();
                    const x = caja.left + caja.width / 2;
                    const y = caja.top + caja.height * 0.62;
                    const elemento = document.elementFromPoint(x, y);
                    if (elemento) {
                        elemento.click();
                        return true;
                    }
                    return false;
                }"""
            )
            if confirmado:
                try:
                    expect(self.page.get_by_text(re.compile(r"Confirmar datos", re.IGNORECASE)).first).not_to_be_visible(timeout=30000)
                except Exception:
                    pass
                self.page.wait_for_timeout(1000)
                if not self._texto_visible_tienda(r"^Confirmar datos$"):
                    return
        except Exception:
            pass

        self.page.mouse.click(942, 526)
        self.page.wait_for_timeout(1000)

    @allure.step("Validar compra realizada en Tienda Usuario")
    def validar_compra_tienda_realizada(self):
        modal_compra = self.page.locator(".modal-content, app-modal, .swal2-popup").filter(
            has_text=re.compile(r"Compra realizada|Gracias por tu compra|Descargar comprobante", re.IGNORECASE)
        ).last
        expect(modal_compra).to_be_visible(timeout=180000)
        self._take_screenshot("tienda_compra_realizada")
        try:
            self.page.get_by_role("button", name=re.compile(r"cerrar|close|×|x", re.IGNORECASE)).first.click(timeout=5000)
        except Exception:
            try:
                self.page.locator(".btn-close, .close, [aria-label='Close'], [aria-label='Cerrar']").first.click(timeout=5000)
            except Exception:
                pass

    def _click_tienda_por_texto(self, texto_regex: str, timeout: int = 30000):
        patron = re.compile(texto_regex, re.IGNORECASE)
        opciones = [
            self.page.get_by_role("button", name=patron).first,
            self.page.get_by_role("link", name=patron).first,
            self.page.get_by_text(patron).first,
            self.page.locator("button, a, ion-button").filter(has_text=patron).first,
        ]
        ultimo_error = None
        for opcion in opciones:
            try:
                if opcion.count() == 0:
                    continue
                opcion.click(timeout=timeout)
                try:
                    self.page.wait_for_load_state("networkidle", timeout=10000)
                except Exception:
                    pass
                self.page.wait_for_timeout(500)
                return
            except Exception as error:
                ultimo_error = error
                continue
        raise AssertionError(f"No se pudo hacer clic en tienda sobre: {texto_regex}. Error: {ultimo_error}")

    @allure.step("Login Corporativo - Código: '{codigo}', Usuario: '{usuario}'")
    def login_corp(self, codigo: str, usuario: str, passw: str):
        self._cerrar_modal_error_conexion()
        campo_codigo = self.page.get_by_placeholder("Ingresar Código")
        if campo_codigo.count() == 0 or not campo_codigo.first.is_visible():
            try:
                self.page.locator("a[href='/login-corporate']").click(timeout=5000)
            except Exception:
                self.page.get_by_role("button", name="Usuario Corporativo").click(timeout=10000)
            self._cerrar_modal_error_conexion()

        test_ids_por_pais = {
            "guatemala": "pw-login-option-country-gt",
            "honduras": "pw-login-option-country-hn",
            "el salvador": "pw-login-option-country-sv",
        }
        country_test_id = test_ids_por_pais.get(self.pais_actual.strip().lower())
        pais_login_corp = self.page.locator(f"[data-testid='{country_test_id}']") if country_test_id else self.page.locator("not-found")
        if pais_login_corp.count() > 0 and pais_login_corp.first.is_visible():
            pais_login_corp.first.click(timeout=10000)
            self._cerrar_modal_error_conexion()

        campo_codigo = self.page.get_by_placeholder("Ingresar Código")
        if campo_codigo.count() == 0 or not campo_codigo.first.is_visible():
            self.page.get_by_role("button", name="Usuario Corporativo").click(timeout=10000)
            self._cerrar_modal_error_conexion()
            campo_codigo = self.page.get_by_placeholder("Ingresar Código")
        expect(campo_codigo).to_be_visible(timeout=30000)
        campo_codigo.click()
        campo_codigo.fill(codigo)
        
        self._fill_app_input("[data-testid='pw-logincorporate-input-username']", usuario)
        self._fill_app_input("[data-testid='pw-logincorporate-input-password']", passw)
        
        self.page.get_by_role("button", name="Iniciar sesión").click()
        self.page.wait_for_url(f"**/historico", wait_until="networkidle")
        self._take_screenshot("login_corp_success")

    def _fill_app_input(self, selector: str, value: str):
        campo = self.page.locator(selector).first
        expect(campo).to_be_visible(timeout=30000)
        input_interno = campo.locator("input").first
        try:
            if input_interno.count() > 0:
                input_interno.click(timeout=5000)
                input_interno.fill(value, timeout=5000)
                return
        except Exception:
            pass

        campo.click()
        self.page.keyboard.press("Control+A")
        self.page.keyboard.type(value)

    def _cerrar_modal_error_conexion(self):
        try:
            modal_error = self.page.get_by_role("dialog", name=re.compile(r"Error de Conexi[oó]n", re.IGNORECASE))
            if modal_error.count() > 0 and modal_error.first.is_visible():
                modal_error.get_by_role("button", name="Aceptar").click(timeout=5000)
                self.page.wait_for_timeout(1000)
        except Exception:
            pass

    @allure.step("Login Express Center - Estación: '{estacion}', Correo: '{correo}'")
    def login_exec(self, estacion: str, correo: str, passw: str):
        self.page.get_by_role("button", name="Usuario Express Center").click()
        self.page.wait_for_timeout(800)

        # En /login-exc a veces el selector de estación aparece directo (Angular state),
        # otras veces la página vuelve a pedir seleccionar país primero.
        # Intento 1: selector de estación directo (flujo estándar)
        btn_sel = self.page.locator("button, ion-button").filter(
            has_text=re.compile(r"Selecciona tu estaci", re.I)
        ).first
        try:
            btn_sel.wait_for(state="visible", timeout=5000)
            btn_sel.click()
        except Exception:
            # Intento 2: /login-exc muestra país primero — seleccionar Guatemala y reintentar
            pais_btn = self.page.get_by_text("Guatemala", exact=True).first
            pais_btn.wait_for(state="visible", timeout=10000)
            pais_btn.click()
            self.page.wait_for_timeout(600)
            self.page.get_by_role("button", name="Selecciona tu estación").click(timeout=15000)

        self.page.get_by_role("searchbox", name="search text").click()
        self.page.get_by_role("searchbox", name="search text").fill(estacion)
        self.page.get_by_text(estacion, exact=True).click()

        self.page.get_by_role("textbox", name="Ingresar Usuario").click()
        self.page.get_by_role("textbox", name="Ingresar Usuario").fill(correo)

        self.page.get_by_role("textbox", name="Ingresar Contraseña").click()
        self.page.get_by_role("textbox", name="Ingresar Contraseña").fill(passw)

        self.page.get_by_role("button", name="Iniciar sesión").click()
        self.page.wait_for_url(f"**/historico", wait_until="networkidle")
        self._take_screenshot("login_exec_success")

    @allure.step("Login App Courier - Teléfono: '{telefono}'")
    def login_courier_app(self, telefono: str, entorno: str = None):
        self.page.wait_for_timeout(1000)
        input_phone = self.page.locator("#inputPhone")
        if input_phone.count() > 0:
            input_phone.fill(telefono)
        else:
            self.page.locator("input:not([type='hidden'])").last.fill(telefono)
        self.page.wait_for_timeout(1000)
        
        self.page.get_by_role("button", name="Autenticar").click()
        self.page.get_by_text("Autenticación").wait_for()
        
        # Use provided entorno or instance entorno
        entorno_a_usar = entorno if entorno is not None else self.entorno
        self.token = self.obtener_token_db(telefono, entorno_a_usar)
        self.ingresar_token(self.token)
        self._take_screenshot("login_courier_success")

    @allure.step("Ingresar Token OTP")
    def ingresar_token(self, token: str):
        if not token or not token.strip():
            raise ValueError("Token inválido o vacío.")
            
        for i, digito in enumerate(token):
            self.page.locator(f"#inputToken{i + 1}").fill(digito)
            
        self.page.get_by_text("Iniciar sesión").click()
        self.page.wait_for_load_state("networkidle")
        self.page.wait_for_timeout(2000)
        self._take_screenshot("token_entered")

    # ==============================================================================
    # ACCIONES DE NAVEGACIÓN INTERNA
    # ==============================================================================

    @allure.step("Seleccionar opción tipo de servicio: '{tipo_servicio}'")
    def opcion_tipo_servicio(self, tipo_servicio: str):
        self.page.get_by_title(tipo_servicio).click()
        self._take_screenshot("service_type_selected")

    @allure.step("Seleccionar opción Crear Guías")
    def opcion_crear_guias(self):
        self.page.get_by_role("link", name="chevron forward Crear Guías").click()
        self.page.wait_for_timeout(1000)
        self._take_screenshot("create_guides_section")

    @allure.step("Seleccionar opción Servicios en EXEC")
    def opcion_servicios_exec(self):
        candidates = [
            self.page.get_by_role("link", name=re.compile(r"^Servicios$", re.I)),
            self.page.get_by_role("button", name=re.compile(r"^Servicios$", re.I)),
            self.page.get_by_role("menuitem", name=re.compile(r"^Servicios$", re.I)),
            self.page.get_by_text("Servicios", exact=True),
            self.page.get_by_text("Servicios", exact=False),
        ]

        for locator in candidates:
            try:
                if locator.count() > 0:
                    locator.first.click(timeout=10000)
                    self.page.wait_for_load_state("networkidle", timeout=20000)
                    self._take_screenshot("servicios_selected_exec")
                    return
            except Exception:
                continue

        raise AssertionError("No se pudo seleccionar la opción 'Servicios' en EXEC")

    @allure.step("Verificar módulo Servicios en EXEC")
    def verificar_modulo_servicios_exec(self):
        validaciones = [
            self.page.get_by_role("searchbox", name="search text"),
            self.page.get_by_text("No se encontraron servicios", exact=False),
            self.page.get_by_text("Servicios", exact=False),
        ]

        for locator in validaciones:
            try:
                if locator.count() > 0 and locator.first.is_visible(timeout=10000):
                    self._take_screenshot("servicios_module_exec")
                    return
            except Exception:
                continue

        current_url = self.page.url.lower()
        if "serv" in current_url:
            self._take_screenshot("servicios_module_exec_url")
            return

        raise AssertionError("No se validó la carga del módulo de Servicios en EXEC")

    @allure.step("Seleccionar opción de servicio en EXEC: '{opcion}'")
    def seleccionar_opcion_servicio_exec(self, opcion: str):
        options = [
            self.page.get_by_role("link", name=opcion),
            self.page.get_by_role("button", name=opcion),
            self.page.get_by_text(opcion, exact=True),
            self.page.get_by_text(opcion, exact=False),
        ]

        for locator in options:
            try:
                if locator.count() > 0:
                    locator.first.click(timeout=10000)
                    self.page.wait_for_load_state("networkidle", timeout=20000)
                    opcion_slug = re.sub(r"[^a-z0-9]+", "_", opcion.lower()).strip("_")
                    self._take_screenshot(f"servicio_{opcion_slug}_selected")
                    return
            except Exception:
                continue

        raise AssertionError(f"No se encontró la opción de servicio '{opcion}'")

    @allure.step("Ingresar guía en recepción EXEC")
    def ingresar_guia_recepcion_exec(self, nombre_textbox: str, guia: str):
        textbox = self.page.get_by_role("textbox", name=nombre_textbox)
        if textbox.count() == 0:
            textbox = self.page.get_by_placeholder(nombre_textbox)
        if textbox.count() == 0:
            textbox = self.page.locator(
                "input[placeholder*='guía' i], input[placeholder*='guia' i], "
                "input[placeholder*='referencia' i], textarea[placeholder*='guía' i], "
                "textarea[placeholder*='referencia' i]"
            )
        if textbox.count() == 0:
            textbox = self.page.locator(
                "input[aria-label*='guía' i], input[aria-label*='guia' i], "
                "input[aria-label*='referencia' i], input[name*='guia' i], "
                "input[name*='referencia' i], input[id*='guia' i], input[id*='referencia' i]"
            )
        if textbox.count() == 0:
            textbox = self.page.locator("input, textarea")
        if textbox.count() == 0:
            raise AssertionError(f"No se encontró el textbox '{nombre_textbox}'")

        textbox.first.wait_for(state="visible", timeout=10000)
        textbox.first.fill(guia, timeout=10000)
        self._take_screenshot("guia_recepcion_ingresada")

    @allure.step("Click en botón EXEC: '{nombre_boton}'")
    def click_boton_exec(self, nombre_boton: str):
        if nombre_boton.strip().lower() == "agregar":
            boton = self.page.get_by_role("button", name=re.compile(r"^Agregar$", re.I))
            if boton.count() == 0:
                boton = self.page.get_by_role("button", name="AGREGAR")
            if boton.count() == 0:
                boton = self.page.get_by_text("AGREGAR", exact=True)
            if boton.count() == 0:
                boton = self.page.locator("button:has-text('AGREGAR')")
            if boton.count() == 0:
                boton = self.page.get_by_text(re.compile(r"^AGREGAR$", re.I))
        elif nombre_boton.strip().lower() == "finalizar":
            boton = self.page.get_by_role("button", name=re.compile(r"checkmark\s*circle\s*outline", re.I))
            if boton.count() == 0:
                boton = self.page.get_by_text(re.compile(r"checkmark\s*circle\s*outline", re.I))
            if boton.count() == 0:
                boton = self.page.get_by_role("button", name=re.compile(r"^finalizar$", re.I))
            if boton.count() == 0:
                boton = self.page.get_by_text("FINALIZAR", exact=True)
            if boton.count() == 0:
                boton = self.page.locator("button:has-text('FINALIZAR')")
            if boton.count() == 0:
                boton = self.page.get_by_text(re.compile(r"^finalizar$", re.I))
            if boton.count() == 0:
                boton = self.page.get_by_title("Finalizar y pagar")
            if boton.count() == 0:
                boton = self.page.get_by_title(re.compile(r"finalizar", re.I))
            if boton.count() == 0:
                boton = self.page.locator("button[title*='Finalizar' i], [role='button'][title*='Finalizar' i]")
            if boton.count() == 0:
                boton = self.page.locator("button, [role='button']").filter(has_text=re.compile(r"finalizar", re.I))
        else:
            boton = self.page.get_by_role("button", name=re.compile(re.escape(nombre_boton), re.I))
            if boton.count() == 0:
                boton = self.page.get_by_text(nombre_boton, exact=False)

        if boton.count() == 0:
            raise AssertionError(f"No se encontró el botón '{nombre_boton}'")

        boton.first.click(timeout=10000)
        self.page.wait_for_timeout(800)
        self._take_screenshot(f"boton_{re.sub(r'[^a-z0-9]', '_', nombre_boton.lower())}_clicked")

    @allure.step("Click en botón CONTINUAR en EXEC")
    def click_boton_continuar_exec(self, nombre_boton: str = "CONTINUAR"):
        boton = self.page.get_by_role("button", name="Continuar chevron forward")
        if boton.count() == 0:
            boton = self.page.get_by_role("button", name=re.compile(r"continuar\s*chevron\s*forward", re.I))
        if boton.count() == 0:
            boton = self.page.get_by_role("button", name=re.compile(r"^continuar$", re.I))
        if boton.count() == 0:
            boton = self.page.get_by_text("CONTINUAR", exact=True)
        if boton.count() == 0:
            boton = self.page.locator("button:has-text('CONTINUAR'), [role='button']:has-text('CONTINUAR')")
        if boton.count() == 0:
            boton = self.page.locator("button, [role='button']").filter(has_text=re.compile(r"continuar", re.I))

        if boton.count() == 0:
            raise AssertionError(f"No se encontró el botón '{nombre_boton}'")

        boton.first.click(timeout=10000)
        self.page.wait_for_timeout(800)
        self._take_screenshot("boton_continuar_clicked")

    @allure.step("Ingresar nombre de cliente en EXEC")
    def ingresar_nombre_cliente_exec(self, nombre_cliente: str):
        textbox = self.page.get_by_role("textbox", name="Nombre de cliente")
        if textbox.count() == 0:
            textbox = self.page.get_by_placeholder("Nombre de cliente")
        if textbox.count() == 0:
            textbox = self.page.locator(
                "input[placeholder*='cliente' i], textarea[placeholder*='cliente' i], "
                "input[aria-label*='cliente' i], input[name*='cliente' i], input[id*='cliente' i]"
            )
        if textbox.count() == 0:
            textbox = self.page.locator("input[type='text'], textarea")
        if textbox.count() == 0:
            raise AssertionError("No se encontró el textbox 'Nombre de cliente'")

        textbox.first.wait_for(state="visible", timeout=10000)
        textbox.first.fill(nombre_cliente, timeout=10000)
        self._take_screenshot("nombre_cliente_ingresado")

    @allure.step("Ingresar DPI en EXEC")
    def ingresar_dpi_exec(self, dpi: str):
        textbox = self.page.get_by_role("textbox", name="DPI")
        if textbox.count() == 0:
            textbox = self.page.get_by_placeholder("DPI")
        if textbox.count() == 0:
            textbox = self.page.locator(
                "input[placeholder*='dpi' i], input[aria-label*='dpi' i], "
                "input[name*='dpi' i], input[id*='dpi' i]"
            )
        if textbox.count() == 0:
            textbox = self.page.locator("input[type='number'], input[type='text']")
        if textbox.count() == 0:
            raise AssertionError("No se encontró el textbox 'DPI'")

        textbox.first.wait_for(state="visible", timeout=10000)
        textbox.first.fill(dpi, timeout=10000)
        self._take_screenshot("dpi_ingresado")

    @allure.step("Verificar link de servicio en EXEC")
    def verificar_link_servicio_exec(self, nombre_link: str):
        link = self.page.get_by_role("link", name=nombre_link)
        if link.count() == 0:
            link = self.page.get_by_text(nombre_link, exact=False)
        if link.count() == 0:
            raise AssertionError(f"No se encontró el link '{nombre_link}'")

        expect(link.first).to_be_visible(timeout=10000)
        link_slug = re.sub(r"[^a-z0-9]+", "_", nombre_link.lower()).strip("_")
        self._take_screenshot(f"link_{link_slug}_visible")

    @allure.step("Verificar link de recepción en EXEC")
    def verificar_link_recepcion_exec(self, nombre_link: str):
        self.verificar_link_servicio_exec(nombre_link)

    @allure.step("Verificar link de devolución en EXEC")
    def verificar_link_devolucion_exec(self, nombre_link: str):
        self.verificar_link_servicio_exec(nombre_link)

    # ==============================================================================
    # FLUJO PRINCIPAL: CREACIÓN DE GUÍAS
    # ==============================================================================

    @allure.step("Creación de {cantidad} guías con parámetros específicos")
    def creacion_guias(self, cantidad: int, direccion): # 'direccion' es el dataclass de test_forza_steps.py
        modal_ya_fue_cerrado = False
        is_collet = str(direccion.collet).lower() == 'true'
        is_nueva_origen = str(direccion.direccion_nueva_origen).lower() == 'true'
        is_nueva_destino = str(direccion.direccion_nueva_destino).lower() == 'true'
        is_std = direccion.tipo_guia == "Servicio Estándar"
        
        for i in range(cantidad):
            print(f"Registrando solicitud #{i + 1}")
            monto_aleatorio = 0
            self.page.wait_for_timeout(1000)
            self.page.get_by_title(direccion.tipo_guia).click()

            if is_nueva_origen:
                # Aquí puedes poner la URL mapeada desde tu configuración
                self.page.wait_for_url("**/opciones-guias", wait_until="networkidle", timeout=30000)
                self.page.get_by_text("Crear dirección de origen").click()
                self.crear_nueva_direccion_origen()

            self.page.locator("#step5 #buttonOpenSelect").click()

            if is_nueva_destino:
                self.page.locator("//strong[contains(text(), 'Crear nueva dirección')]").click()
                self.page.wait_for_timeout(1000)
                self.crear_nueva_direccion_destino()
            else:
                self.page.get_by_text(direccion.nombre_direccion).click()
                self.page.wait_for_timeout(1000)

            if is_collet:
                self.page.get_by_label("Cobro contra entrega").check()
            else:
                self.page.get_by_label("Cobro a mi cuenta").check()

            self.page.locator("#step8 input[type='text']").fill("QA002")
            self.page.wait_for_timeout(1000)

            if is_std:
                self.page.get_by_text("Agregar paquete").click()
            else:
                monto_aleatorio = random.randint(5000, 20000)
                self.page.locator("#step19 input[type='text']").fill(str(monto_aleatorio))
                self.page.locator("#step20 #buttonOpenSelect").click()
                # ¡AQUÍ ESTÁ LA MAGIA! Reemplazamos el string quemado por direccion.tarjeta
                self.page.get_by_text(direccion.tarjeta).click()
                self.page.get_by_text("Agregar paquete").click()
                self.page.wait_for_timeout(1000)

            if not modal_ya_fue_cerrado:
                self.page.locator("#packageInformationModal button").click()
                modal_ya_fue_cerrado = True

            self.page.wait_for_timeout(1000)
            self.page.get_by_title("Finalizar y pagar").click()
            self.page.wait_for_timeout(1000)
            self.page.get_by_role("button", name="Si").click()
            self._take_screenshot(f"guide_{i+1}_confirm")

            if monto_aleatorio > 10000:
                self.page.wait_for_timeout(1000)
                self.page.get_by_role("button", name="Si").click()

            if is_collet:
                self.page.get_by_text("Finalizar", exact=True).click()
                self.page.locator("div").filter(has_text=re.compile(r"^¿Cómo realizarás tu recolección\?$")).locator("button").click()
                self.page.get_by_title("Mis Envíos").click()
                self.page.wait_for_timeout(1000)
                
                numero_guia = self.page.locator("(//table//tr[1]//strong[contains(text(), 'FD')])[1]").inner_text()
                print(f"Número de guía: {numero_guia}")
                
                if is_std:
                    self.guardar_guia_en_archivo("guias_Collet_STD", "STD", numero_guia)
                else:
                    self.guardar_guia_en_archivo("guias_Collet_COD", "COD", numero_guia)
                self._take_screenshot(f"guide_{i+1}_created")
            else:
                self.page.wait_for_timeout(1000)
                self.page.get_by_label("CF").check()
                self.page.get_by_label("Tarjeta").check()
                self.page.get_by_text("Pagar Carrito").click()
                self.page.wait_for_timeout(1000)
                self.page.wait_for_load_state("load")

                frame = self.page.frame_locator("#RequestFAC")
                texto_frame = frame.get_by_role("paragraph").inner_text()

                match = re.search(r"FD\d+", texto_frame)
                if match:
                    guia = match.group(0)
                    if is_std:
                        self.guardar_guia_en_archivo("guias_STD_Prepagada", "STD", guia)
                    else:
                        self.guardar_guia_en_archivo("guias_Collet_COD_Prepagada", "COD", guia)
                    print(f"Guía encontrada: {guia}")
                else:
                    print("No se encontró una guía.")
                self._take_screenshot(f"guide_{i+1}_prepaid")

                # TRADUCCIÓN DEL PATRÓN "Task" DE C# A PYTHON (Eventos asíncronos en API Sync)
                with self.page.expect_popup() as popup_info, self.page.expect_download() as download_info:
                    self.page.locator("#confirmPaymentModal button").click()
                
                popup = popup_info.value
                download = download_info.value

                self.page.locator("div").filter(has_text=re.compile(r"^¿Cómo realizarás tu recolección\?$")).locator("button").click()

    # ==============================================================================
    # DIRECCIONES Y AYUDANTES DE FORMULARIOS
    # ==============================================================================

    @allure.step("Llenar formulario de nueva dirección destino")
    def crear_nueva_direccion_destino(self):
        nombre = self.generar_nombre_qa()
        telefono = self.generar_telefono()

        self.page.get_by_placeholder("Ingresar Nombre").nth(1).click()
        self.page.get_by_placeholder("Ingresar Nombre").nth(1).fill(nombre)

        self.page.get_by_placeholder("Lugar en la ciudad").nth(1).click()
        self.page.locator("//li[normalize-space(.)='Gimnasio']").nth(1).click()

        self.page.get_by_placeholder("Seleccione Poblado").nth(1).click()
        self.page.locator("//li[normalize-space(.)='Cahabon, Santa Maria Cahabon, Alta Verapaz']").nth(1).click()

        self.page.get_by_placeholder("Ingresar Teléfono").nth(1).click()
        self.page.get_by_placeholder("Ingresar Teléfono").nth(1).fill(telefono)

        self.page.get_by_placeholder("Ingrese la dirección exacta").nth(1).click()
        self.page.get_by_placeholder("Ingrese la dirección exacta").nth(1).fill("Cahabón, Guatemala")

        self.page.locator("//input[@placeholder='Ingrese el nombre de la persona que recibirá el envío' and @type='text']").click()
        self.page.locator("//input[@placeholder='Ingrese el nombre de la persona que recibirá el envío' and @type='text']").fill("QA Demo")

        # Forzar validaciones (Presionar espacio)
        self.page.get_by_placeholder("Ingresar Nombre").nth(1).press(" ")
        self.page.get_by_placeholder("Ingresar Teléfono").nth(1).press(" ")
        self.page.locator("//input[@placeholder='Ingrese el nombre de la persona que recibirá el envío' and @type='text']").press(" ")
        self.page.get_by_placeholder("Ingrese la dirección exacta").nth(1).press(" ")

        self.page.locator("//button[normalize-space(.)='Guardar']").nth(3).click()
        self._take_screenshot("new_destination_address")

    @allure.step("Llenar formulario de nueva dirección origen")
    def crear_nueva_direccion_origen(self):
        nombre = self.generar_nombre_qa()
        telefono = self.generar_telefono()

        self.page.get_by_placeholder("Ingresar Nombre").nth(0).click()
        self.page.get_by_placeholder("Ingresar Nombre").nth(0).fill(nombre)

        self.page.get_by_placeholder("Lugar en la ciudad").nth(0).click()
        self.page.locator("//li[normalize-space(.)='Gimnasio']").nth(0).click()

        self.page.get_by_placeholder("Seleccione Poblado").nth(0).click()
        self.page.locator("//li[normalize-space(.)='Cahabon, Santa Maria Cahabon, Alta Verapaz']").nth(0).click()

        self.page.get_by_placeholder("Ingresar Teléfono").nth(0).click()
        self.page.get_by_placeholder("Ingresar Teléfono").nth(0).fill(telefono)

        self.page.get_by_placeholder("Ingrese la dirección exacta").nth(0).click()
        self.page.get_by_placeholder("Ingrese la dirección exacta").nth(0).fill("Cahabón, Guatemala")

        self.page.get_by_placeholder("Ingrese ubicación ej. zona 10").click()
        self.page.get_by_placeholder("Ingrese ubicación ej. zona 10").fill("Ciudad de guatemala")

        input_ub = self.page.get_by_placeholder("Ingrese ubicación ej. zona 10")
        input_ub.press("Backspace")

        self.page.get_by_text("Ciudad de Guatemala", exact=True).click()

        self.page.get_by_placeholder("Ingresar Nombre").nth(0).press(" ")
        self.page.get_by_placeholder("Ingresar Teléfono").nth(0).press(" ")
        self.page.locator("//input[@placeholder='Ingrese el nombre de la persona que recibirá el envío' and @type='text']").press(" ")
        self.page.get_by_placeholder("Ingrese la dirección exacta").nth(0).press(" ")

        self.page.locator(".form-group > .btn").first.click()
        self._take_screenshot("new_origin_address")

    # ==============================================================================
    # PROCESAMIENTO DE EXCEL (Equivalente a ClosedXML)
    # ==============================================================================

    def ruta_excel_obtener(self, ruta: str):
        self.ruta_excel = ruta

    def hoja_excel(self, hoja: str):
        self.hoja = hoja

    def columna_excel(self, col: str):
        self.columna = col

    @allure.step("Ingresar lote de recolección desde Excel (Inicio: {inicial}, Final: {final}, Piezas: {cantidad_piezas})")
    def ingresar_a_recoleccion(self, inicial: int, final: int, cantidad_piezas: int = 1):
        datos = self.obtener_guias_desde_excel(inicial, final)
        self.elegir_recolectar()

        textbox = self.page.locator("//*[@id='lblGuide']")
        for valor in datos:
            textbox.fill(f"{str(valor)}-{cantidad_piezas}")
            with self.page.expect_response("**/Home.aspx/FDApi", timeout=60000):
                textbox.press("Enter")
            textbox.wait_for(state="visible")
        self._take_screenshot("batch_collection_complete")

        # Pasos finales para completar la recolección
        self.page.get_by_role("button", name="Siguiente").click()
        
        # Llenar correo del cliente
        correo_textbox = self.page.get_by_role("textbox", name="Correo del cliente *")
        correo_textbox.wait_for(state="visible", timeout=30000)
        correo_textbox.fill("carlos.fernandez@forzalatam.com")
        
        # Dibujar en el canvas para activar botón Finalizar
        canvas = self.page.locator("canvas").first
        canvas.wait_for(state="visible")
        
        # Inyectar firma en el canvas para activar el botón Finalizar
        canvas = self.page.locator("canvas").first
        canvas.wait_for(state="visible")
        
        # Dibujar en el canvas con eventos reales (lo que funcionó visualmente)
        bbox = canvas.bounding_box()
        if bbox:
            self.page.mouse.move(bbox['x'] + 50, bbox['y'] + 100)
            self.page.mouse.down()
            for i in range(10):
                self.page.mouse.move(bbox['x'] + 50 + i * 15, bbox['y'] + 100 + (i % 3) * 10)
            self.page.mouse.up()
            self.page.wait_for_timeout(500)
        
        # Estrategia: Forzar el botón y ejecutar su acción directamente
        self.page.evaluate("""
            () => {
                const btn = document.querySelector('#btnFinish');
                if (!btn) return;
                
                // Habilitar el botón visualmente
                btn.disabled = false;
                btn.removeAttribute('disabled');
                btn.classList.remove('disabled');
                
                // Intentar obtener la función que ejecuta el botón
                // 1. Revisar si tiene un onclick directo
                if (btn.onclick) {
                    console.log('Ejecutando onclick directo');
                    btn.onclick();
                    return;
                }
            
                // 3. Si no encuentra nada, disparar un click event estándar
                console.log('Disparando evento click estándar');
                btn.dispatchEvent(new MouseEvent('click', {
                    view: window,
                    bubbles: true,
                    cancelable: true
                }));
            }
        """)
        
        # Esperar el mensaje de éxito
        self.page.get_by_role("alert").wait_for(state="visible", timeout=30000)
        expect(self.page.get_by_role("alert")).to_contain_text("Servicio completado exitosamente.")
        self._take_screenshot("recoleccion_finalizada")

    def elegir_recolectar(self):
        self.page.locator("//button[@type='button' and @data-type='Pickup' and contains(text(), 'Recolectar')]").first.click()
        self.page.locator("//button[@type='button' and contains(text(), 'Recolectar')]").click()

    @allure.step("Abrir recolección pendiente en posición {posicion}")
    def abrir_recoleccion_por_posicion(self, posicion: int):
        if posicion < 1:
            raise ValueError("La posición de la recolección debe ser mayor o igual a 1.")

        botones_recolectar = self.page.locator("//button[contains(normalize-space(.), 'Recolectar')]")
        botones_recolectar.nth(posicion - 1).click(timeout=30000)
        self.page.wait_for_load_state("networkidle")
        self._take_screenshot(f"pickup_position_{posicion}")

    @allure.step("Reportar visita fallida")
    def reportar_visita_fallida(self, ruta_imagen: str):
        self.page.get_by_role("button", name=re.compile("Visita fallida", re.I)).click(timeout=30000)
        self.page.wait_for_load_state("networkidle")

        self.seleccionar_primer_motivo_visita_fallida()
        self.cargar_imagen_visita_fallida(ruta_imagen)

        expect(self.page.get_by_text(re.compile("Ubicación detectada exitosamente", re.I))).to_be_visible(timeout=30000)
        reportar = self.page.locator("//button[contains(normalize-space(.), 'Reportar')]").last
        reportar.scroll_into_view_if_needed(timeout=10000)
        expect(reportar).to_be_enabled(timeout=120000)
        self.page.wait_for_function(
            """
            () => {
                const buttons = [...document.querySelectorAll('button')];
                const reportar = buttons.reverse().find((button) => button.textContent.includes('Reportar'));
                if (!reportar) return false;
                const style = window.getComputedStyle(reportar);
                return !reportar.disabled && style.visibility !== 'hidden' && style.display !== 'none';
            }
            """,
            timeout=120000,
        )
        reportar.click(timeout=120000)
        self._take_screenshot("failed_pickup_reported")

    def seleccionar_primer_motivo_visita_fallida(self):
        selected = self.page.evaluate(
            """
            () => {
                const select = document.querySelector('select');
                if (!select) return false;
                if (select.options.length > 1) {
                    select.selectedIndex = 1;
                }
                select.dispatchEvent(new Event('input', { bubbles: true }));
                select.dispatchEvent(new Event('change', { bubbles: true }));
                return Boolean(select.value);
            }
            """
        )
        if selected:
            return

        motivo_default = self.page.get_by_text("Información incompleta", exact=True)
        if motivo_default.count() > 0:
            return

        comboboxes = self.page.get_by_role("combobox")
        if comboboxes.count() > 0:
            try:
                comboboxes.first.select_option(index=1)
                return
            except Exception:
                pass

        selectores = self.page.locator("select")
        if selectores.count() > 0:
            opciones = selectores.first.locator("option")
            if opciones.count() > 1:
                selectores.first.select_option(index=1)
            else:
                selectores.first.select_option(index=0)
            return

        try:
            self.page.get_by_role("combobox").first.click(timeout=10000)
            self.page.get_by_role("option").nth(1).click(timeout=10000)
        except Exception:
            return

    def cargar_imagen_visita_fallida(self, ruta_imagen: str):
        ruta = Path(ruta_imagen)
        if not ruta.is_absolute():
            ruta = Path(__file__).resolve().parent.parent / ruta
        if not ruta.exists():
            raise FileNotFoundError(f"No existe la imagen para visita fallida: {ruta}")

        file_input = self.page.locator("input[type='file']").first
        file_input.set_input_files(str(ruta))
        self.page.wait_for_timeout(1000)

    @allure.step("Validar mensaje de visita fallida: '{mensaje}'")
    def validar_mensaje_visita_fallida(self, mensaje: str):
        expect(self.page.get_by_text(mensaje)).to_be_visible(timeout=120000)
        self._take_screenshot("failed_pickup_success_message")

    def obtener_guias_desde_excel(self, inicial: int, final: int) -> list:
        datos = self.leer_columna_excel(self.ruta_excel, self.hoja, self.columna)
        if inicial < 1:
            raise ValueError("El rango inicial debe ser mayor o igual a 1.")
        if final < inicial:
            raise ValueError("El rango final debe ser mayor o igual al rango inicial.")

        limite = min(len(datos), final)
        return datos[inicial - 1:limite]

    def leer_columna_excel(self, nombre_archivo: str, hoja: str, columna_letra: str) -> list:
        ruta_absoluta = os.path.abspath(nombre_archivo)
        if not os.path.exists(ruta_absoluta):
            raise FileNotFoundError(f"El archivo no existe: {ruta_absoluta}")

        try:
            wb = openpyxl.load_workbook(ruta_absoluta, data_only=True)
            if hoja not in wb.sheetnames:
                raise ValueError(f"La hoja '{hoja}' no existe.")
            
            ws = wb[hoja]
            col_idx = column_index_from_string(columna_letra)
            valores = []

            for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, values_only=True):
                if row[0] is not None:
                    valores.append(row[0])
            return valores
        except Exception as e:
            print(f"Error leyendo el archivo Excel: {e}")
            return []

    # ==============================================================================
    # INTEGRACIÓN DE BASE DE DATOS (SQL Server con pyodbc)
    # ==============================================================================
    
    @allure.step("Obtener token OTP desde la Base de Datos")
    def obtener_token_db(self, telefono: str, entorno: str = None) -> str:
        """
        Conecta a SQL Server dinámicamente según el entorno, busca el token más reciente
        para un teléfono y lo devuelve.
        """
        entorno_normalizado = (entorno or "DEFAULT").upper()
        env_key = f"SQLSERVER_{entorno_normalizado}_CONNECTION_STRING"
        base_connection_string = os.getenv(env_key) or os.getenv("SQLSERVER_CONNECTION_STRING")

        if not base_connection_string:
            raise ValueError(
                f"No se encontró cadena de conexión. Configura {env_key} "
                "o SQLSERVER_CONNECTION_STRING en el archivo .env."
            )

        full_connection_string = base_connection_string.rstrip(";")
        if "Encrypt=" not in full_connection_string:
            full_connection_string += ";Encrypt=yes"
        if "TrustServerCertificate=" not in full_connection_string:
            full_connection_string += ";TrustServerCertificate=yes"
        full_connection_string += ";"

        try:
            # 4. Abrir la conexión
            conn = pyodbc.connect(full_connection_string)
            cursor = conn.cursor()

            # 5. La consulta SQL
            # IMPORTANTE: En pyodbc no se usan nombres como '@Telefono', se usan signos de interrogación '?'
            query = """
                SELECT TOP 1 srt.LoginToken
                FROM dbo.SenderReceiver sr
                INNER JOIN dbo.SenderReceiverLoginToken srt 
                    ON sr.ID = srt.SenderReceiverId
                WHERE sr.Phone = ?
                ORDER BY srt.DateCreated DESC
            """

            # 6. Equivalente a command.Parameters.Add() y ExecuteScalarAsync()
            # Pasamos 'telefono' como una tupla al método execute para prevenir SQL Injection
            cursor.execute(query, (telefono,))
            row = cursor.fetchone()

            # 7. Equivalente a result?.ToString()
            return str(row[0]) if row else ""

        except pyodbc.Error as e:
            print(f"Error de base de datos conectando a SQL Server: {e}")
            return ""
        except Exception as e:
            print(f"Error general en la obtención del token: {e}")
            return ""
        finally:
            # Buena práctica: Asegurarnos de cerrar la conexión en Python equivalente al 'await using'
            if 'conn' in locals():
                conn.close()

    # ==============================================================================
    # GENERADORES DE DATOS Y ARCHIVOS UTILS
    # ==============================================================================

    def guardar_guia_en_archivo(self, nombre_archivo: str, tipo_guia: str, guia: str):
        if not nombre_archivo.endswith(".txt"):
            nombre_archivo += ".txt"
            
        linea = f"Guia {tipo_guia} ----> {guia}\n"
        with open(nombre_archivo, "a", encoding="utf-8") as f:
            f.write(linea)

    def generar_nombre_qa(self) -> str:
        nombres = ["Carlos", "Ana", "Luis", "Sofía", "Pedro", "Laura", "Mario", "Paola", "Alfita", "Juancito", "Manuelito"]
        nombre_aleatorio = random.choice(nombres)
        numero = random.randint(100, 999)
        return f"QA {nombre_aleatorio}{numero}"

    def generar_telefono(self) -> str:
        primer_digito = str(random.randint(3, 7))
        resto = "".join(str(random.randint(0, 9)) for _ in range(7))
        return primer_digito + resto

    def generar_cuenta_aleatoria(self, longitud: int) -> str:
        return "".join(str(random.randint(0, 9)) for _ in range(longitud))

    # ==============================================================================
    # PAGO ZIGI (Paggo)
    # ==============================================================================

    @allure.step("Completar pago Zigi en link de pago")
    def completar_pago_zigi(self, link: str, nombre: str, apellido: str, email: str,
                             telefono: str, tarjeta: str, vencimiento: str, cvv: str,
                             dpi: str, departamento: str, direccion: str):
        self.page.goto(link, timeout=120000)
        self.page.wait_for_load_state("load")
        self.page.wait_for_timeout(1000)
        self._take_screenshot("zigi_pagina_pago_cargada")

        self.page.locator('input[name="nombre"]').wait_for(state="visible", timeout=60000)
        self._take_screenshot("zigi_formulario_visible")

        self.page.locator('input[name="nombre"]').click()
        self.page.locator('input[name="nombre"]').fill(nombre)
        self.page.locator('input[name="apellido"]').click()
        self.page.locator('input[name="apellido"]').fill(apellido)
        self.page.locator('input[name="email-field"]').click()
        self.page.locator('input[name="email-field"]').fill(email)
        self.page.locator('input[name="phone-field"]').click()
        self.page.locator('input[name="phone-field"]').fill(telefono)
        self._take_screenshot("zigi_datos_personales")
        self.page.wait_for_timeout(1000)

        self.page.locator('input[autocomplete="cc-number"]').click()
        self.page.locator('input[autocomplete="cc-number"]').fill(tarjeta)
        self.page.locator('input[autocomplete="cc-exp"]').click()
        self.page.locator('input[autocomplete="cc-exp"]').fill(vencimiento)
        self.page.locator('input[placeholder="123"]').click()
        self.page.locator('input[placeholder="123"]').fill(cvv)
        self._take_screenshot("zigi_datos_tarjeta")
        self.page.wait_for_timeout(1000)

        self.page.get_by_role("textbox", name="Ingrese su DPI o NIT").click()
        self.page.get_by_role("textbox", name="Ingrese su DPI o NIT").fill(dpi)
        self.page.get_by_role("button", name="Departamento*").first.click()
        self.page.get_by_role("option", name=departamento).click()
        self.page.wait_for_timeout(500)
        self.page.get_by_role("button", name=f"Departamento* {departamento}").nth(1).click()
        self.page.get_by_role("option", name=departamento).click()
        self.page.locator('input[name="direccionUno"]').click()
        self.page.locator('input[name="direccionUno"]').fill(direccion)
        self._take_screenshot("zigi_datos_facturacion")
        self.page.wait_for_timeout(1000)

        self.page.route("https://www.google.com/**", lambda route: route.abort())

        self.page.locator('[data-test-id="btn-loading-button"]').dblclick()

        self.page.get_by_text("Transaccion exitosa").wait_for(state="visible", timeout=60000)
        self._take_screenshot("zigi_transaccion_exitosa")

        self.page.get_by_text("Pago realizado exitosamente").wait_for(state="visible", timeout=10000)
        self._take_screenshot("zigi_comprobante_pago")

        print(f"  ✔ Pago Zigi completado — Transacción exitosa")
        allure.attach(
            f"Link: {link}\nNombre: {nombre} {apellido}\nResultado: Transaccion exitosa",
            name="Pago Zigi Completado",
            attachment_type=allure.attachment_type.TEXT,
        )

    # ==============================================================================
    # MIS ENVÍOS - PORTAL EXPRESS CENTER (EXEC)
    # ==============================================================================

    def _esperar_mis_envios_lista(self, timeout: int = 30000):
        self.page.wait_for_load_state("networkidle", timeout=timeout)
        self.page.get_by_role("searchbox", name="search text").wait_for(state="visible", timeout=timeout)

    def _guias_en_tabla(self):
        return self.page.locator("ion-label").filter(has_text=re.compile(r"FD\d+"))

    @allure.step("Verificar que la pantalla Mis Envíos cargó correctamente")
    def verificar_mis_envios_cargada(self):
        self._esperar_mis_envios_lista()
        guias = self._guias_en_tabla()
        sin_resultados = self.page.get_by_text("No se encontraron servicios", exact=False)
        try:
            guias.first.wait_for(state="visible", timeout=10000)
        except Exception:
            assert sin_resultados.is_visible(), \
                "Mis Envíos cargó pero no muestra guías FD (ion-label) ni mensaje de estado vacío"
        self._take_screenshot("mis_envios_cargada")

    @allure.step("Aplicar filtro de estado: '{estado}'")
    def filtrar_mis_envios_por_estado(self, estado: str):
        btn = self.page.locator("ion-segment-button").filter(has_text=estado).first
        btn.wait_for(state="visible", timeout=15000)
        btn.click()
        self.page.wait_for_load_state("networkidle", timeout=20000)
        self.page.wait_for_timeout(1000)
        self._take_screenshot(f"filtro_{re.sub(r'[^a-z0-9]', '_', estado.lower())}")

    @allure.step("Verificar que la tabla refleja el filtro: '{estado}'")
    def verificar_filtro_mis_envios(self, estado: str):
        self.page.wait_for_load_state("networkidle", timeout=10000)
        guias = self._guias_en_tabla()
        try:
            guias.first.wait_for(state="visible", timeout=3000)
        except Exception:
            pass
        self._take_screenshot(f"verificacion_{re.sub(r'[^a-z0-9]', '_', estado.lower())}")

    @allure.step("Obtener número de la primera guía en Mis Envíos")
    def obtener_primera_guia_mis_envios(self) -> str:
        self._esperar_mis_envios_lista()
        guia_locator = self._guias_en_tabla().first
        guia_locator.wait_for(state="visible", timeout=30000)
        guia = guia_locator.inner_text().strip()
        print(f"Primera guía encontrada: {guia}")
        self._take_screenshot("primera_guia_obtenida")
        return guia

    @allure.step("Buscar en Mis Envíos por texto: '{texto}'")
    def buscar_en_mis_envios(self, texto: str):
        search_input = self.page.get_by_role("searchbox", name="search text")
        search_input.wait_for(state="visible", timeout=15000)
        search_input.fill("")
        search_input.fill(texto)
        self.page.wait_for_timeout(1000)
        self._take_screenshot(f"busqueda_{texto[:20].replace('/', '_')}")

    @allure.step("Verificar que la tabla muestra únicamente la guía: '{guia}'")
    def verificar_resultado_busqueda_guia(self, guia: str):
        self.page.wait_for_timeout(500)
        guias = self._guias_en_tabla()
        assert guias.count() == 1, f"Se esperaba 1 resultado, se encontraron {guias.count()}"
        expect(guias.first).to_contain_text(guia)
        self._take_screenshot("verificacion_busqueda_guia")

    @allure.step("Verificar que la tabla no muestra resultados")
    def verificar_sin_resultados_mis_envios(self):
        self.page.wait_for_timeout(2000)
        invalidas = self.page.locator("ion-label").filter(
            has_text=re.compile(r"GUIA-NO-EXISTE", re.I)
        )
        assert invalidas.count() == 0, \
            "Apareció un resultado inesperado para el texto de búsqueda inválido"
        self._take_screenshot("sin_resultados_mis_envios")

    @allure.step("Activar filtro por fecha y buscar desde '{fecha_inicio}' hasta '{fecha_fin}'")
    def filtrar_mis_envios_por_fecha(self, fecha_inicio: str, fecha_fin: str):
        from datetime import date as _date
        toggle = self.page.locator("label").filter(has_text=re.compile(r"Filtrar por Fecha", re.I)).first
        toggle.wait_for(state="visible", timeout=10000)
        toggle.click()
        self.page.wait_for_timeout(500)

        d_s, m_s, y_s = (int(x) for x in fecha_inicio.split("/"))
        d_e, m_e, y_e = (int(x) for x in fecha_fin.split("/"))
        hoy = _date.today()

        cal_icons = self.page.get_by_role("img", name="Calendar")
        cal_icons.first.click()
        self.page.wait_for_timeout(400)

        meses_back = (hoy.year - y_s) * 12 + (hoy.month - m_s)
        btn_prev = self.page.get_by_role("button", name="‹")
        for _ in range(meses_back):
            if btn_prev.first.get_attribute("disabled") is not None:
                raise AssertionError(
                    f"El calendario no permite navegar hasta {fecha_inicio} — "
                    "error preexistente: rango mínimo del datepicker limitado"
                )
            btn_prev.first.click()
            self.page.wait_for_timeout(150)

        self.page.locator("span").filter(has_text=re.compile(rf"^{d_s}$")).first.click()
        self.page.wait_for_timeout(300)

        cal_icons.nth(1).click()
        self.page.wait_for_timeout(400)

        meses_back_end = (hoy.year - y_e) * 12 + (hoy.month - m_e)
        for _ in range(meses_back_end):
            btn_prev.first.click()
            self.page.wait_for_timeout(150)

        self.page.locator("span").filter(has_text=re.compile(rf"^{d_e}$")).first.click()
        self.page.wait_for_timeout(300)

        self.page.get_by_text("Buscar guías").click()
        self.page.wait_for_load_state("networkidle", timeout=30000)
        self._take_screenshot("filtro_por_fecha_aplicado")

    @allure.step("Verificar que la tabla se actualizó con el filtro de fechas")
    def verificar_filtro_fecha_mis_envios(self):
        self.page.wait_for_timeout(1000)
        guias = self._guias_en_tabla()
        sin_resultados = self.page.get_by_text("No se encontraron servicios", exact=False)
        assert guias.count() > 0 or sin_resultados.is_visible(), \
            "La tabla no mostró resultados ni mensaje vacío tras el filtro de fechas"
        self._take_screenshot("verificacion_filtro_fecha")

    @allure.step("Verificar paginación y navegar a la siguiente página si existe")
    def verificar_paginacion_mis_envios(self):
        self.page.wait_for_load_state("networkidle", timeout=15000)
        btn_p2 = self.page.get_by_text("page 2").first
        if btn_p2.count() > 0:
            btn_p2.click()
            self.page.wait_for_load_state("networkidle", timeout=20000)
            guias_p2 = self._guias_en_tabla()
            try:
                guias_p2.first.wait_for(state="visible", timeout=15000)
                assert guias_p2.count() > 0, "La página 2 no cargó ninguna guía"
            except Exception:
                sin_resultados = self.page.get_by_text("No se encontraron servicios", exact=False)
                assert sin_resultados.is_visible(), "La página 2 no cargó guías ni mensaje de estado"
            self._take_screenshot("paginacion_pagina_2")
        else:
            print("Solo hay una página de resultados — paginación verificada sin navegación")
            self._take_screenshot("paginacion_unica_pagina")

    @allure.step("Clic en el número de la primera guía para abrir rastreo")
    def abrir_rastreo_primera_guia_mis_envios(self):
        guia_label = self._guias_en_tabla().first
        guia_label.wait_for(state="visible", timeout=30000)
        guia_label.click()
        self.page.wait_for_timeout(2000)
        self._take_screenshot("rastreo_abierto")

    @allure.step("Verificar que el rastreo se abrió con información de la guía")
    def verificar_rastreo_mis_envios(self):
        rastreo = self.page.locator(".modal.show, [role='dialog'], app-rastreo, app-modal, .modal-dialog").first
        guia_visible = self.page.locator("//strong[contains(text(),'FD')] | //h4[contains(text(),'FD')] | //h5[contains(text(),'FD')]").first
        assert rastreo.is_visible(timeout=10000) or guia_visible.is_visible(timeout=10000), \
            "No se abrió el rastreo ni apareció el número de guía en pantalla"
        self._take_screenshot("rastreo_verificado")

    def _encontrar_elemento_visible_por_texto(self, texto: str, patron: str = None) -> object:
        busqueda = re.compile(patron or texto, re.I)
        for tag in ("label", "ion-label", "span", "button"):
            candidatos = self.page.locator(tag).filter(has_text=busqueda)
            for i in range(candidatos.count()):
                el = candidatos.nth(i)
                if el.is_visible():
                    return el
        return None

    def _diagnostico_dom(self):
        resultado = self.page.evaluate("""
            () => {
                const tags = ['button', 'label', 'ion-label', 'span', 'a'];
                const out = [];
                for (const tag of tags) {
                    for (const el of document.querySelectorAll(tag)) {
                        const txt = el.textContent.trim().replace(/\\s+/g, ' ');
                        if (txt && el.offsetParent !== null) {
                            out.push('<' + tag + '> "' + txt + '"');
                        }
                    }
                }
                return out;
            }
        """)
        print("\n=== DOM DIAGNOSIS: elementos visibles ===")
        for t in resultado[:60]:
            try:
                print(t)
            except UnicodeEncodeError:
                print(t.encode("ascii", "replace").decode("ascii"))
        print("=== FIN DIAGNÓSTICO ===\n")

    @allure.step("Verificar que el botón Generar Excel está visible y habilitado")
    def verificar_boton_generar_excel(self):
        self._esperar_mis_envios_lista()
        self.page.wait_for_timeout(500)
        el = self._encontrar_elemento_visible_por_texto("Generar Excel")
        if el is None:
            self._diagnostico_dom()
        assert el is not None, \
            "No se encontró ningún elemento visible con texto 'Generar Excel' en label/ion-label/span/button"
        self._take_screenshot("boton_generar_excel_visible")

    @allure.step("Completar flujo Generar Excel: seleccionar Marzo 2026 y descargar")
    def abrir_modal_generar_excel(self):
        el = self._encontrar_elemento_visible_por_texto("Generar Excel")
        assert el is not None, "No se encontró 'Generar Excel' al intentar hacer clic"
        el.click()
        self.page.wait_for_timeout(2000)
        self._take_screenshot("modal_generar_excel_abierto")

        self.page.get_by_role("textbox", name="Seleccione un Año").click()
        self.page.wait_for_timeout(600)
        self._take_screenshot("modal_picker_anio_abierto")
        self.page.get_by_role("heading", name="2026").click()
        self.page.wait_for_timeout(500)

        self.page.get_by_role("textbox", name=re.compile(r"Seleccione un mes", re.I)).click()
        self.page.wait_for_timeout(600)
        self._take_screenshot("modal_picker_mes_abierto")
        self.page.get_by_role("heading", name="Marzo").click()
        self.page.wait_for_timeout(500)
        self._take_screenshot("modal_generar_excel_seleccionado")

        with self.page.expect_download(timeout=20000) as dl_info:
            self.page.get_by_role("button", name=re.compile(r"Aceptar", re.I)).click()

        download = dl_info.value
        assert download.suggested_filename != "", "La descarga del Excel no se completó"
        self._take_screenshot("excel_descargado")
        print(f"[OK] Excel descargado: {download.suggested_filename}")

    @allure.step("Filtrar Envíos completados, seleccionar guías y descargar PDF")
    def verificar_boton_reimprimir_guias(self):
        self._esperar_mis_envios_lista()

        completados = self.page.locator("ion-segment-button").filter(
            has_text=re.compile(r"completados", re.I)
        ).first
        completados.wait_for(state="visible", timeout=10000)
        completados.click()
        self.page.wait_for_load_state("networkidle", timeout=20000)
        self.page.wait_for_timeout(1000)
        self._take_screenshot("reimprimir_filtro_completados")

        checkboxes = self.page.get_by_role("checkbox")
        checkboxes.nth(3).wait_for(state="visible", timeout=15000)
        checkboxes.nth(3).click()
        self.page.wait_for_timeout(800)
        self._take_screenshot("reimprimir_checkbox_seleccionado")

        with self.page.expect_download(timeout=20000) as dl_info:
            self.page.get_by_text("Generar PDF").nth(1).click()

        download = dl_info.value
        assert download.suggested_filename != "", "La descarga del PDF no se completó"
        self._take_screenshot("pdf_descargado")
        print(f"[OK] PDF descargado: {download.suggested_filename}")

    # ==============================================================================
    # FLUJOS CORPORATIVOS Y EXEC
    # ==============================================================================

    @allure.step("Crear guías CORPORATIVO")
    def crear_guias_corp(self, datos):
        is_std = datos.tipo_guia == "Servicio Estándar"
        is_collet = str(datos.collet).lower() == "true"

        self.page.get_by_role("textbox", name="Selecciona un poblado").click()
        self.page.get_by_role("heading", name="Cahabon, Santa Maria Cahabon").click()
        self.page.get_by_role("button", name="calculator outline Calcular").click()

        if is_std:
            card = self.page.locator("div.card").filter(has_text="Servicio Estándar")
            card.get_by_role("button", name="Seleccionar").click()
        else:
            self.page.get_by_role("button", name="Seleccionar send").nth(1).click()

        self.page.get_by_role("radio").nth(1).click()
        self.page.get_by_role("textbox", name="Ingresar quién recibe").click()
        self.page.get_by_role("textbox", name="Ingresar quién recibe").fill("Juan")

        self.page.locator("input[name='ion-input-18']").click()
        self.page.locator("input[name='ion-input-18']").fill("Juan")
        self.page.locator("input[name='ion-input-19']").click()
        self.page.locator("input[name='ion-input-19']").fill("52452421")

        self.page.get_by_role("textbox", name="Ingresar dirección en destinatario").click()
        self.page.get_by_role("textbox", name="Ingresar dirección en destinatario").fill("Ciudad de Guatemala")
        self.page.get_by_role("button", name="Siguiente arrow forward").click()

        if datos.tipo_guia == "Servicio C.O.D.":
            self.page.get_by_role("textbox", name="Ingresar monto COD").click()
            self.page.get_by_role("textbox", name="Ingresar monto COD").fill("100")

        self.page.get_by_role("button", name="Siguiente arrow forward").click()

        if is_collet:
            self.page.locator("ion-col").filter(has_text=re.compile(r"^Crédito$")).locator("label").click()
            self.page.locator("ion-col").filter(has_text=re.compile(r"^Collect$")).locator("label").click()

        self.page.get_by_text("Arrow Back Regresar Mostrar").click()
        self.page.get_by_role("button", name="Mostrar Resumen arrow forward").click()
        self.page.wait_for_timeout(1000)

        heading_locator = self.page.locator("role=heading")
        matched_heading = heading_locator.filter(has_text="No. Guía Transporte FD").first

        if matched_heading.count() > 0:
            full_text = matched_heading.text_content()
            match = re.search(r"FD\d+", full_text)
            if match:
                guia = match.group(0)
                if is_std:
                    nombre_archivo = "guias_Corp_Collet_STD" if is_collet else "guias_Corp_Credito_STD"
                    tipo = "STD Collet" if is_collet else "STD Credito"
                else:
                    nombre_archivo = "guias_Corp_Collet_COD" if is_collet else "guias_Corp_Credito_COD"
                    tipo = "COD Collet" if is_collet else "COD Credito"
                print(f"Guía encontrada: {guia}")
                self.guardar_guia_en_archivo(nombre_archivo, tipo, guia)
                self._take_screenshot("corp_guide_created")
        else:
            print("No se encontró el heading con la guía.")

    @allure.step("Crear guía CORPORATIVO - Servicio de devolución")
    def crear_guia_devolucion_corp(self, datos):
        self.page.locator("label").filter(has_text="¿Es una devolución?").click()
        try:
            expect(self.page.get_by_text("El servicio será una devolución")).to_be_visible(timeout=5000)
        except Exception:
            self.page.locator("ion-toggle").first.click()
            expect(self.page.get_by_text("El servicio será una devolución")).to_be_visible(timeout=30000)

        self._seleccionar_destino_devolucion(datos)
        self.page.get_by_role("button", name=re.compile(r"Calcular", re.IGNORECASE)).click()
        self._seleccionar_servicio_devolucion(datos.tipo_guia)

        expect(self.page.get_by_text("Remitente", exact=True)).to_be_visible(timeout=30000)
        if self._es_destino_manual(datos.destino):
            self.page.get_by_placeholder("Ingresar nombre de contacto").first.fill(datos.nombre_contacto)
            self.page.get_by_placeholder("Ingresar teléfono").first.fill(datos.telefono)
            self.page.get_by_role("textbox", name="Ingresar dirección en remitente").fill(datos.direccion_remitente)

        self._seleccionar_tipo_entrega(datos.tipo_entrega)
        self.page.get_by_role("button", name=re.compile(r"Siguiente", re.IGNORECASE)).click()

        expect(self.page.get_by_text("Cobro de Servicio")).to_be_visible(timeout=30000)
        self._seleccionar_tipo_cobro(datos.tipo_cobro)
        self.page.get_by_role("button", name=re.compile(r"Mostrar Resumen", re.IGNORECASE)).click()

        self.page.wait_for_function("() => /No\\. Guía:\\s*FD\\d+/.test(document.body.innerText)", timeout=60000)
        resumen = self.page.locator("body").inner_text()
        match = re.search(r"No\.\s*Guía:\s*(FD\d+)", resumen)
        if not match:
            raise AssertionError("No se encontró el número de guía de devolución en el resumen.")

        self.guia_generada = match.group(1)
        self._validar_resumen_devolucion(datos, resumen)

        pais = self._normalizar_nombre_archivo(datos.pais)
        destino = self._normalizar_nombre_archivo(datos.destino)
        servicio = "COD" if "cod" in datos.tipo_guia.lower() else "STD"
        cobro = self._normalizar_nombre_archivo(datos.tipo_cobro)
        self.guardar_guia_en_archivo(f"guias_Corp_Devolucion_{pais}_{destino}_{servicio}_{cobro}", f"{servicio} Devolucion {datos.tipo_cobro}", self.guia_generada)
        print(f"Guía de devolución encontrada: {self.guia_generada}")
        self._take_screenshot("corp_return_guide_created")

    def _seleccionar_destino_devolucion(self, datos):
        destino = datos.destino.strip().lower()
        if self._es_destino_manual(datos.destino):
            self._seleccionar_opcion_destino("Manual")
            self._seleccionar_poblado_devolucion(datos.poblado)
            return

        if "cartera" in destino:
            self._seleccionar_opcion_destino("Cartera de Clientes")
        elif "exc" in destino or "express center" in destino or "libreta" in destino:
            self._seleccionar_opcion_destino("Libreta Express Center")
        else:
            raise AssertionError(f"Destino de devolución no soportado: {datos.destino}")

        if datos.referencia_destino:
            self._seleccionar_referencia_destino(datos.referencia_destino)

    def _seleccionar_referencia_destino(self, referencia_destino: str):
        expect(self.page.get_by_text(referencia_destino).first).to_be_visible(timeout=30000)
        contenedores = [
            "ion-card",
            "ion-item",
            "ion-row",
            ".card",
            ".container",
        ]

        for selector in contenedores:
            contenedor = self.page.locator(selector).filter(has_text=re.compile(re.escape(referencia_destino), re.IGNORECASE))
            if contenedor.count() == 0:
                continue

            objetivos = []
            radio = contenedor.first.locator("ion-radio")
            if radio.count() > 0:
                objetivos.append(radio.first)
            boton_interno = contenedor.first.locator("button")
            if boton_interno.count() > 0:
                objetivos.append(boton_interno.first)
            objetivos.append(contenedor.first)

            for objetivo in objetivos:
                try:
                    objetivo.click(timeout=5000)
                    if self._esperar_modal_seleccion_cerrado():
                        return
                except Exception:
                    continue

        self.page.get_by_text(referencia_destino).first.click(timeout=10000)
        if not self._esperar_modal_seleccion_cerrado():
            raise AssertionError(f"No se pudo confirmar la selección del destino: {referencia_destino}")

    def _esperar_modal_seleccion_cerrado(self) -> bool:
        try:
            expect(self.page.locator("ion-modal.show-modal")).not_to_be_visible(timeout=10000)
            return True
        except Exception:
            return False

    def _seleccionar_poblado_devolucion(self, poblado: str):
        campo_poblado = self.page.get_by_role("textbox", name="Selecciona un poblado")
        campo_poblado.click()
        candidatos = [
            poblado,
            self._quitar_acentos(poblado),
            poblado.split(",", 1)[0],
            self._quitar_acentos(poblado.split(",", 1)[0]),
        ]

        for candidato in candidatos:
            try:
                campo_poblado.fill(candidato, timeout=5000)
            except Exception:
                pass

            try:
                self.page.get_by_role("heading", name=candidato).click(timeout=5000)
                return
            except Exception:
                pass

            try:
                self.page.locator("role=heading").filter(has_text=re.compile(re.escape(candidato), re.IGNORECASE)).first.click(timeout=5000)
                return
            except Exception:
                continue

        raise AssertionError(f"No se pudo seleccionar el poblado de devolución: {poblado}")

    def _seleccionar_opcion_destino(self, opcion: str):
        try:
            self.page.get_by_role("radio", name=re.compile(opcion, re.IGNORECASE)).click(timeout=5000)
            return
        except Exception:
            pass

        try:
            self.page.get_by_text(opcion, exact=True).click(timeout=10000)
        except Exception:
            self.page.locator("ion-radio").filter(has_text=re.compile(opcion, re.IGNORECASE)).click(timeout=10000)

    def _seleccionar_servicio_devolucion(self, tipo_guia: str):
        servicio = "Servicio C.O.D." if "cod" in tipo_guia.lower() else "Servicio Estándar"
        card = self.page.locator("div.card").filter(has_text=servicio)
        card.get_by_role("button", name="Seleccionar").click()

    def _seleccionar_tipo_entrega(self, tipo_entrega: str):
        try:
            self.page.get_by_role("radio", name=re.compile(tipo_entrega, re.IGNORECASE)).click(timeout=5000)
        except Exception:
            indice = 1 if "oficina" in tipo_entrega.lower() else 0
            self.page.locator("ion-radio:visible").nth(indice).click(timeout=10000)

    def _seleccionar_tipo_cobro(self, tipo_cobro: str):
        if "collect" not in tipo_cobro.lower():
            return

        try:
            self.page.locator("ion-col").filter(has_text=re.compile(r"^Crédito$")).locator("label").click(timeout=5000)
        except Exception:
            pass

        try:
            self.page.locator("ion-col").filter(has_text=re.compile(r"^Collect$")).locator("label").click(timeout=10000)
        except Exception:
            self.page.get_by_text("Collect", exact=True).click(timeout=10000)

    def _validar_resumen_devolucion(self, datos, resumen: str):
        valores_esperados = [self.guia_generada]
        if self._es_destino_manual(datos.destino):
            valores_esperados.extend([datos.nombre_contacto, datos.telefono, datos.direccion_remitente])
        elif datos.referencia_destino:
            valores_esperados.append(datos.referencia_destino)

        for valor in valores_esperados:
            if valor and valor not in resumen:
                raise AssertionError(f"El resumen de devolución no contiene el dato esperado: {valor}")

    def _es_destino_manual(self, destino: str) -> bool:
        return "manual" in destino.strip().lower()

    def _normalizar_nombre_archivo(self, value: str) -> str:
        value = value.strip().lower()
        value = self._quitar_acentos(value)
        return re.sub(r"[^a-z0-9]+", "_", value).strip("_")

    def _quitar_acentos(self, value: str) -> str:
        reemplazos = {
            "á": "a",
            "é": "e",
            "í": "i",
            "ó": "o",
            "ú": "u",
            "ñ": "n",
        }
        for origen, destino in reemplazos.items():
            value = value.replace(origen, destino)
            value = value.replace(origen.upper(), destino.upper())
        return value

    @allure.step("Validar guía de devolución en Mis Envíos")
    def validar_guia_devolucion_en_mis_envios(self):
        if not self.guia_generada:
            raise AssertionError("No hay una guía de devolución guardada para validar.")

        self._ir_a_mis_envios_corporativo()
        guia = self.page.get_by_text(self.guia_generada, exact=True).first
        try:
            expect(guia).to_be_visible(timeout=120000)
        except Exception:
            self.page.reload(wait_until="networkidle", timeout=120000)
            expect(guia).to_be_visible(timeout=120000)
        self._take_screenshot("corp_return_guide_in_my_shipments")

    def _ir_a_mis_envios_corporativo(self):
        patron_mis_envios = re.compile(r"MIS ENV[IÍ]OS|Mis Env[ií]os", re.IGNORECASE)
        opciones = [
            self.page.get_by_role("button", name=patron_mis_envios).first,
            self.page.get_by_role("link", name=patron_mis_envios).first,
            self.page.get_by_text(patron_mis_envios).first,
        ]

        for opcion in opciones:
            try:
                if opcion.count() == 0:
                    continue
                opcion.click(timeout=10000)
                self.page.wait_for_load_state("networkidle", timeout=30000)
                return
            except Exception:
                continue

        self._navegar_a_historico_corporativo()

    def _navegar_a_historico_corporativo(self):
        match = re.match(r"^(https?://[^/]+)", self.page.url or self.url_actual)
        if not match:
            raise AssertionError("No se pudo resolver la URL base para abrir Mis Envíos.")

        base_url = match.group(1)
        for ruta in ("/express/historico", "/historico"):
            try:
                self.page.goto(f"{base_url}{ruta}", wait_until="networkidle", timeout=120000)
                if "historico" in self.page.url:
                    return
            except Exception:
                continue

        raise AssertionError("No se pudo abrir Mis Envíos para validar la guía generada.")

    @allure.step("Crear guías EXEC")
    def crear_guias_exec(self, datos):
        is_std = datos.tipo_guia == "Servicio Estándar"
        is_collet = str(datos.collet).lower() == "true"

        self.page.locator("ion-radio").filter(has_text="Clientes de Express Center").click()
        self.page.get_by_role("heading", name="Use el buscador para ver").click()
        self.page.wait_for_timeout(1000)
        self.page.get_by_role("textbox", name="Selecciona un poblado").click()
        self.page.get_by_role("heading", name="Cahabon, Santa Maria Cahabon").click()
        self.page.get_by_role("button", name="calculator outline Calcular").click()
        self.page.wait_for_timeout(1000)

        if is_std:
            card = self.page.locator("div.card").filter(has_text="Servicio Estándar")
            card.get_by_role("button", name="Seleccionar").click()
        else:
            card = self.page.locator("div.card").filter(has_text="Servicio C.O.D.")
            card.get_by_role("button", name="Seleccionar").click()

        nombre_envia = self.generar_nombre_qa()
        telefono_envia = self.generar_telefono()

        self.page.get_by_placeholder("Ingresar nombre de contacto").first.click()
        self.page.get_by_placeholder("Ingresar nombre de contacto").first.fill(nombre_envia)
        self.page.get_by_placeholder("Ingresar teléfono").first.click()
        self.page.get_by_placeholder("Ingresar teléfono").first.fill(telefono_envia)
        self.page.get_by_role("textbox", name="Ingresar correo electrónico").first.click()
        self.page.get_by_role("textbox", name="Ingresar correo electrónico").first.fill("carlos.fernandez@forzalatam.com")

        nombre_recibe = self.generar_nombre_qa()
        self.page.get_by_placeholder("Ingresar nombre de contacto").nth(1).click()
        self.page.get_by_placeholder("Ingresar nombre de contacto").nth(1).fill(nombre_recibe)

        telefono_recibe = self.generar_telefono()
        self.page.get_by_placeholder("Ingresar teléfono").nth(1).click()
        self.page.get_by_placeholder("Ingresar teléfono").nth(1).fill(telefono_recibe)

        self.page.get_by_role("textbox", name="Ingresar dirección en destinatario").click()
        self.page.get_by_role("textbox", name="Ingresar dirección en destinatario").fill("Ciudad de Guatemala")

        self.page.get_by_role("radio").nth(1).click()
        self.page.get_by_role("button", name="Siguiente arrow forward").click()

        if datos.tipo_guia == "Servicio C.O.D.":
            self.page.get_by_role("textbox", name="Ingresar monto COD").click()
            self.page.get_by_role("textbox", name="Ingresar monto COD").fill("100")
            self.page.get_by_role("textbox", name="Seleccione Banco").click()
            self.page.get_by_role("heading", name="BANRURAL - BANCO DE").click()
            self.page.get_by_role("textbox", name="Seleccione Tipo de Cuenta").click()
            self.page.get_by_role("heading", name="Ahorro").click()
            self.page.get_by_role("textbox", name="Ingresar número de cuenta").click()
            cuenta = self.generar_cuenta_aleatoria(15)
            self.page.get_by_role("textbox", name="Ingresar número de cuenta").fill(cuenta)
            self.page.get_by_role("textbox", name="Ingresar nombre de la cuenta").click()
            self.page.get_by_role("textbox", name="Ingresar nombre de la cuenta").fill("QA")
            self.page.get_by_role("textbox", name="Ingresar Documento de").click()
            self.page.get_by_role("textbox", name="Ingresar Documento de").fill("6352452")

        self.page.get_by_role("button", name="Siguiente arrow forward").click()

        if is_collet:
            self.page.locator("ion-col").filter(has_text=re.compile(r"^Collect$")).locator("label").click()

        self.page.get_by_text("Arrow Back Regresar Mostrar").click()
        self.page.get_by_role("button", name="Mostrar Resumen arrow forward").click()
        self.page.wait_for_timeout(1000)

        heading_locator = self.page.locator("role=heading")
        matched_heading = heading_locator.filter(has_text="No. Guía Transporte FD").first

        if matched_heading.count() > 0:
            full_text = matched_heading.text_content()
            match = re.search(r"FD\d+", full_text)
            if match:
                guia = match.group(0)
                if is_std:
                    nombre_archivo = "guias_EXEC_Collet_STD" if is_collet else "guias_Corp_Credito_STD"
                    tipo = "STD Collet" if is_collet else "STD Credito"
                else:
                    nombre_archivo = "guias_EXEC_Collet_COD" if is_collet else "guias_Corp_Credito_COD"
                    tipo = "COD Collet" if is_collet else "COD Credito"
                print(f"Guía encontrada: {guia}")
                self.guardar_guia_en_archivo(nombre_archivo, tipo, guia)
                self._take_screenshot("exec_guide_created")
        else:
            print("No se encontró el heading con la guía.")
